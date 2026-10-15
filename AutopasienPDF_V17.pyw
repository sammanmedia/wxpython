# -*- coding: utf-8 -*-
"""
Auto List Pasien dari PDF  (v17)
Database : AutopasienPDF.db (SQLite, bisa dibuka pakai DBeaver)
Library  : wxPython, pymupdf, openpyxl, rapidocr-onnxruntime (OCR)
           pip install wxPython pymupdf openpyxl rapidocr-onnxruntime
"""
import difflib
import json
import os
import re
import shutil
import sqlite3
import sys
import threading
import time
import webbrowser
from datetime import datetime
from pathlib import Path

import wx
import wx.grid as gridlib

try:
    import pymupdf as fitz  # nama baru
except ImportError:
    try:
        import fitz  # nama lama
    except ImportError:
        fitz = None
try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None
try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None

# ---- OCR (untuk halaman hasil scan gambar)
try:
    import numpy as np
except Exception:
    np = None
try:
    from rapidocr_onnxruntime import RapidOCR
except Exception:
    RapidOCR = None
try:
    import pytesseract  # cadangan, butuh program Tesseract terpasang
    from PIL import Image
    _tess = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    if os.path.isfile(_tess):
        pytesseract.pytesseract.tesseract_cmd = _tess
except Exception:
    pytesseract = None
    Image = None

OCR_AVAILABLE = np is not None and (RapidOCR is not None or pytesseract is not None)
OCR_DPI = 200
OCR_LOCK = threading.Lock()
_OCR = {"engine": None}


class Cancelled(Exception):
    """Dilempar saat user menekan BATAL PROSES (dipakai untuk menghentikan scan seketika)."""


def get_app_dir():
    """Folder aplikasi: folder .exe kalau sudah di-convert, folder script kalau jalan biasa."""
    if getattr(sys, "frozen", False):
        return os.path.dirname(os.path.abspath(sys.executable))
    return os.path.dirname(os.path.abspath(__file__))


APP_TITLE = "Auto List Pasien dari pdf v17"
DB_NAME = "AutopasienPDF.db"
BASE_DIR = get_app_dir()
DEFAULT_DB = os.path.join(BASE_DIR, DB_NAME)
CONFIG_FILE = os.path.join(BASE_DIR, "AutoListPasien.json")

GENDERS = ["Laki-laki", "Perempuan"]
ROW_HEIGHT = 26  # tinggi baris tetap (tidak bisa ditarik)
MAX_LOG_LIST = 300  # batas baris daftar konflik / data kosong yang ditulis ke log

# ------------------------------------------------------------ struktur kolom
# (nama kolom db, label, tipe SQL)
FIELDS = [
    ("no_SEP", "NO SEP", "TEXT"),
    ("no_peserta", "NO. PESERTA", "TEXT"),
    ("no_rekammedis", "NO REKAM MEDIS", "TEXT"),
    ("jenis_kelamin", "JENIS KELAMIN", "TEXT"),
    ("tgl_masuk", "TGL MASUK", "TEXT"),
    ("tgl_keluar", "TGL KELUAR", "TEXT"),
    ("nama_peserta", "NAMA PESERTA", "TEXT"),
    ("control_ulang", "KONTROL ULANG", "TEXT"),
    ("no_telepon", "NO TELEPON", "TEXT"),
    ("dokter", "DOKTER", "TEXT"),
    ("ceklis", "CEKLIS", "INTEGER DEFAULT 0"),  # 0 = belum, 1 = sudah
]
DB_FIELDS = [f[0] for f in FIELDS]
SEARCH_FIELDS = [k for k in DB_FIELDS if k != "ceklis"]
OBSOLETE_COLS = {"umur", "tanggal_lahir"}  # kolom lama yang dihapus

LABELS = {k: lab for k, lab, _ in FIELDS}
LABELS["id"] = "ID"
LABELS["wa"] = "WHATSAPP"

# Urutan kolom di tabel layar
GRID_KEYS = [
    "id", "no_SEP", "no_peserta", "no_rekammedis", "jenis_kelamin",
    "tgl_masuk", "tgl_keluar", "dokter", "nama_peserta", "control_ulang",
    "no_telepon", "wa", "ceklis",
]
GRID_COLS = [(k, LABELS[k]) for k in GRID_KEYS]
PHONE_COL = GRID_KEYS.index("no_telepon")
WA_COL = GRID_KEYS.index("wa")
CEK_COL = GRID_KEYS.index("ceklis")

COL_WIDTH = {
    "id": 55, "no_SEP": 240, "no_peserta": 140, "no_rekammedis": 120,
    "jenis_kelamin": 120, "tgl_masuk": 100, "tgl_keluar": 100,
    "dokter": 260, "nama_peserta": 230, "control_ulang": 150,
    "no_telepon": 130, "wa": 110, "ceklis": 110,
}

# Kolom untuk export Excel (sama dengan layar, tanpa WhatsApp)
EXPORT_KEYS = [k for k in GRID_KEYS if k != "wa"]

NUMERIC_COLS = {"id", "ceklis"}
DATE_COLS = {"tgl_masuk", "tgl_keluar", "control_ulang"}

COLOR_WHITE = wx.Colour(255, 255, 255)
COLOR_BLACK = wx.Colour(0, 0, 0)
COLOR_GREY = wx.Colour(232, 232, 232)
COLOR_WA = wx.Colour(37, 211, 102)
COLOR_DONE = wx.Colour(198, 239, 206)       # hijau tipis (baris sudah selesai)
COLOR_DONE_TEXT = wx.Colour(0, 97, 0)

# ----------------------------------------------------------- bulan Indonesia
MONTH_FULL = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
              "Agustus", "September", "Oktober", "November", "Desember"]
MONTH_MAP = {n.lower(): i for i, n in enumerate(MONTH_FULL, 1)}
MONTH_MAP.update({
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "jun": 6, "jul": 7, "agu": 8,
    "agt": 8, "ags": 8, "agus": 8, "sep": 9, "sept": 9, "okt": 10, "nov": 11,
    "des": 12,
    # cadangan bahasa Inggris
    "january": 1, "february": 2, "march": 3, "may": 5, "june": 6, "july": 7,
    "august": 8, "october": 10, "december": 12, "aug": 8, "oct": 10, "dec": 12,
})


def month_number(word):
    """Nama bulan -> 1..12. Toleran salah baca OCR (misal 'Agustu5')."""
    w = (word or "").lower().strip(". ")
    if w in MONTH_MAP:
        return MONTH_MAP[w]
    close = difflib.get_close_matches(w, list(MONTH_MAP), n=1, cutoff=0.75)
    return MONTH_MAP[close[0]] if close else None


# ------------------------------------------------------ simpan pengaturan
def load_config():
    try:
        with open(CONFIG_FILE, "r", encoding="utf-8") as f:
            cfg = json.load(f)
        return cfg if isinstance(cfg, dict) else {}
    except Exception:
        return {}


def save_config(**kw):
    cfg = load_config()
    cfg.update(kw)
    try:
        with open(CONFIG_FILE, "w", encoding="utf-8") as f:
            json.dump(cfg, f, ensure_ascii=False, indent=2)
    except Exception:
        pass


def load_last_db():
    """Ambil database terakhir yang dipilih user (kalau file-nya masih ada)."""
    path = load_config().get("db_path", "")
    if path and os.path.isfile(path):
        return path
    return DEFAULT_DB


def save_last_db(path):
    save_config(db_path=path)


# ------------------------------------------------------------------------ OCR
def boxes_to_text(result):
    """Susun hasil RapidOCR (kotak + teks) jadi teks per baris, urut atas ke bawah."""
    if not result:
        return ""
    items = []
    for box, txt, *_ in result:
        ys = [p[1] for p in box]
        xs = [p[0] for p in box]
        items.append((sum(ys) / len(ys), min(xs), max(ys) - min(ys), txt))
    items.sort(key=lambda t: t[0])
    lines = []
    for yc, x, h, txt in items:
        if lines and abs(yc - lines[-1]["y"]) <= 0.6 * max(h, lines[-1]["h"]):
            lines[-1]["items"].append((x, txt))
        else:
            lines.append({"y": yc, "h": h, "items": [(x, txt)]})
    return "\n".join(" ".join(t for _, t in sorted(ln["items"])) for ln in lines)


def run_ocr(rgb):
    """rgb = array numpy (tinggi, lebar, 3) warna RGB. Kembalikan teks hasil OCR."""
    with OCR_LOCK:
        if RapidOCR is not None:
            if _OCR["engine"] is None:
                _OCR["engine"] = RapidOCR()
            bgr = np.ascontiguousarray(rgb[:, :, ::-1])
            result, _ = _OCR["engine"](bgr)
            return boxes_to_text(result)
        if pytesseract is not None:
            return pytesseract.image_to_string(Image.fromarray(rgb), lang="eng")
    return ""


# ---------------------------------------------------------------- PDF parsing
FIELD_LABELS = {
    "no_peserta": ["Nomor Peserta", "No. Peserta", "No Peserta"],
    "no_rekammedis": ["Nomor Rekam Medis", "No. Rekam Medis", "No Rekam Medis", "No. RM"],
    "no_SEP": ["Nomor SEP", "No. SEP", "No SEP"],
    "tgl_masuk": ["Tanggal Masuk", "Tgl. Masuk", "Tgl Masuk"],
    "tgl_keluar": ["Tanggal Keluar", "Tgl. Keluar", "Tgl Keluar"],
}
NAME_LABELS = ["Nama Peserta", "Nama Pasien", "Nama"]
GENDER_LABELS = ["Jenis Kelamin"]

# "No. Telepon : 083134676447" (nilai boleh turun ke baris bawah label)
PHONE_LABEL_RE = re.compile(
    r"No\.?\s*(?:Telepon|Telpon|Telp|Tlp|HP|Handphone)\s*[:：]?\s*(\+?\d[\d\-\. ]{7,22}\d)",
    re.IGNORECASE,
)
# cadangan: nomor HP berpola 08xxxxxxxxxx / 628xxxxxxxxxx / +628xxxxxxxxxx
PHONE_LIKE_RE = re.compile(r"(?<![\d])((?:\+?62|0)8\d{8,12})(?![\d])")

# "Kontrol Ulang : 29 Agustus 2026" (toleran: Kontrol/Control, salah baca OCR)
CONTROL_LABEL_RE = re.compile(r"[KC][o0]ntr[o0][l1I|]\s*U[l1I|]ang", re.IGNORECASE)
DATE_TEXT_RE = re.compile(r"(\d{1,2})\s*([A-Za-z]{3,10})\.?\s*,?\s*(20\d{2})")
DATE_NUM_RE = re.compile(r"(\d{1,2})\s*[/\-.]\s*(\d{1,2})\s*[/\-.]\s*(20\d{2})")

# "Dokter : Dr. Johan H, Sp.PD"  (halaman SURAT ELEGIBILITAS PESERTA / halaman 2)
DOCTOR_RE = re.compile(r"(?<![A-Za-z])Dok[a-z1|]{1,3}r\s*[:：;]\s*([^\n]*)", re.IGNORECASE)
LABEL_START_RE = re.compile(
    r"(?i)^(nama|no|nomor|tanggal|tgl|poli|alamat|jenis|umur|diagnosa|kelas|"
    r"peserta|faskes|kontrol|control|catatan)\b"
)


class PdfDoc:
    def __init__(self, path, stop=None):
        self.doc = None
        self.reader = None
        self.ocr_used = False
        self.stop = stop  # threading.Event, kalau di-set -> proses dihentikan
        if fitz:
            self.doc = fitz.open(path)
            self.count = self.doc.page_count
        elif PdfReader:
            self.reader = PdfReader(path)
            self.count = len(self.reader.pages)
        else:
            raise RuntimeError("Install dulu: pip install pymupdf")
        self._cache = {}
        self._ocr_cache = {}

    def check(self):
        if self.stop is not None and self.stop.is_set():
            raise Cancelled()

    def text(self, i):
        self.check()
        if i in self._cache:
            return self._cache[i]
        try:
            if self.doc is not None:
                t = self.doc[i].get_text("text") or ""
            else:
                t = self.reader.pages[i].extract_text() or ""
        except Exception:
            t = ""
        self._cache[i] = t
        return t

    def is_low_text(self, i):
        """Halaman yang hampir tidak punya teks = kemungkinan hasil scan gambar."""
        return len(self.text(i).strip()) < 60

    def ocr_text(self, i):
        """Render halaman jadi gambar lalu baca dengan OCR (hasil disimpan sementara)."""
        self.check()
        if i in self._ocr_cache:
            return self._ocr_cache[i]
        t = ""
        if self.doc is not None and OCR_AVAILABLE:
            try:
                pix = self.doc[i].get_pixmap(dpi=OCR_DPI, alpha=False)
                arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
                if pix.n == 1:
                    arr = np.repeat(arr, 3, axis=2)
                elif pix.n > 3:
                    arr = arr[:, :, :3]
                t = run_ocr(arr)
                self.ocr_used = True
            except Exception:
                t = ""
        self._ocr_cache[i] = t
        return t

    def close(self):
        if self.doc is not None:
            self.doc.close()


def find_field(text, labels):
    """Cari 'Label : nilai' di teks. Kalau nilai kosong, ambil baris berikutnya."""
    for lab in labels:
        words = lab.split()
        pat = r"(?<![A-Za-z])" + r"\s*".join(re.escape(w) for w in words) + r"\s*[:：]\s*([^\n]*)"
        m = re.search(pat, text, re.IGNORECASE)
        if not m:
            continue
        val = m.group(1).strip()
        if not val:
            rest = text[m.end():].split("\n")
            for line in rest:
                line = line.strip()
                if line:
                    if ":" not in line:
                        val = line
                    break
        val = re.split(r"\s{2,}|\t", val)[0].strip()
        if val:
            return val
    return ""


def normalize_gender(val):
    low = (val or "").lower()
    if "laki" in low:
        return "Laki-laki"
    if "perempuan" in low or "wanita" in low:
        return "Perempuan"
    return ""


def clean_phone(val):
    """Ambil angka saja: '0831-3467-6447' -> '083134676447'."""
    return re.sub(r"\D", "", (val or "").strip())


def phone_from_text(text, exclude):
    """Cari nomor telepon di satu halaman. exclude = nomor lain yang bukan telepon."""
    for m in PHONE_LABEL_RE.finditer(text):
        num = clean_phone(m.group(1))
        if 9 <= len(num) <= 15 and num not in exclude:
            return num
    for m in PHONE_LIKE_RE.finditer(text):
        num = clean_phone(m.group(1))
        if num not in exclude:
            return num
    return ""


def extract_phone(pdf, exclude):
    """
    Urutan pencarian:
    1. halaman yang ada tulisan SURAT ELEGIBILITAS PESERTA
    2. halaman ke-2
    3. halaman ke-1
    4. semua halaman
    """
    order = []
    for i in range(pdf.count):
        t = pdf.text(i).upper()
        if "ELEGIBILITAS" in t or "ELIGIBILITAS" in t:
            order.append(i)
    if pdf.count >= 2 and 1 not in order:
        order.append(1)
    if 0 not in order:
        order.append(0)
    for i in range(pdf.count):
        if i not in order:
            order.append(i)

    for i in order:
        num = phone_from_text(pdf.text(i), exclude)
        if num:
            return num
    return ""


def control_from_text(text):
    """Cari tanggal setelah tulisan 'Kontrol/Control Ulang' (teks PDF atau hasil OCR)."""
    for m in CONTROL_LABEL_RE.finditer(text):
        win = text[m.end(): m.end() + 60]
        cands = []
        for dm in DATE_TEXT_RE.finditer(win):
            mo = month_number(dm.group(2))
            if mo:
                cands.append((dm.start(), f"{dm.group(1)} {MONTH_FULL[mo - 1]} {dm.group(3)}"))
        for dm in DATE_NUM_RE.finditer(win):
            cands.append((dm.start(), f"{dm.group(1)}/{dm.group(2)}/{dm.group(3)}"))
        if cands:
            cands.sort(key=lambda c: c[0])
            return cands[0][1]
    return ""


def doctor_from_text(text):
    """Cari 'Dokter : Dr. Johan H, Sp.PD' -> 'Dr. Johan H, Sp.PD' (teks PDF atau hasil OCR)."""
    for m in DOCTOR_RE.finditer(text):
        val = m.group(1).strip()
        if not val:
            # nilai turun ke baris berikutnya
            for line in text[m.end():].split("\n"):
                line = line.strip()
                if line:
                    val = line
                    break
        val = re.split(r"\s{3,}|\t", val)[0]
        val = re.sub(r"\s+", " ", val).strip(" :;|_-")
        if len(val) < 3 or ":" in val or LABEL_START_RE.match(val):
            continue
        return val
    return ""


def search_pages(pdf, parse, title_words, default_page):
    """
    Cari data di PDF:
    1. teks PDF: halaman berjudul title_words, lalu halaman default, lalu semua halaman
    2. kalau belum ketemu -> OCR: halaman default, lalu halaman yang teksnya sedikit (hasil scan)
    """
    order = []
    for i in range(pdf.count):
        flat = re.sub(r"\s+", " ", pdf.text(i).upper())
        if any(w in flat for w in title_words):
            order.append(i)
    if pdf.count > default_page and default_page not in order:
        order.append(default_page)
    for i in range(pdf.count):
        if i not in order:
            order.append(i)

    for i in order:
        val = parse(pdf.text(i))
        if val:
            return val

    if not OCR_AVAILABLE:
        return ""
    ocr_pages = []
    if pdf.count > default_page:
        ocr_pages.append(default_page)
    for i in order:
        if i not in ocr_pages and pdf.is_low_text(i):
            ocr_pages.append(i)
    for i in ocr_pages:
        val = parse(pdf.ocr_text(i))
        if val:
            return val
    return ""


def extract_control(pdf):
    """Kontrol Ulang: halaman 'SURAT BUKTI PELAYANAN KESEHATAN (SBPK)' / halaman 3."""
    return search_pages(
        pdf, control_from_text,
        ["SURAT BUKTI PELAYANAN KESEHATAN", "SBPK"], 2,
    )


def extract_doctor(pdf, allow_ocr=True):
    """
    Dokter: HANYA dari halaman 'SURAT ELEGIBILITAS PESERTA' / halaman ke-2.
    1. baca teks PDF di halaman tsb
    2. kalau tidak ketemu (misal halaman berupa gambar) -> OCR halaman tsb
    """
    pages = []
    for i in range(pdf.count):
        flat = re.sub(r"\s+", " ", pdf.text(i).upper())
        if "ELEGIBILITAS" in flat or "ELIGIBILITAS" in flat:
            pages.append(i)
    if pdf.count >= 2 and 1 not in pages:
        pages.append(1)

    for i in pages:
        val = doctor_from_text(pdf.text(i))
        if val:
            return val

    if allow_ocr and OCR_AVAILABLE:
        for i in pages:
            val = doctor_from_text(pdf.ocr_text(i))
            if val:
                return val
    return ""


def extract_patient(path, db, stop=None):
    """
    Baca satu PDF. Return (data, existing).
    existing = baris yang sudah ada di database (atau None).
    OCR hanya dijalankan kalau datanya memang belum ada di database.
    Kalau stop di-set, proses berhenti dengan melempar Cancelled.
    """
    pdf = PdfDoc(path, stop)
    try:
        if pdf.count == 0:
            return None, None
        p1 = pdf.text(0)
        data = {k: find_field(p1, labs) for k, labs in FIELD_LABELS.items()}

        # Jenis kelamin (halaman 1 dulu, kalau tidak ada cari semua halaman)
        gender = normalize_gender(find_field(p1, GENDER_LABELS))
        if not gender:
            for i in range(pdf.count):
                gender = normalize_gender(find_field(pdf.text(i), GENDER_LABELS))
                if gender:
                    break
        data["jenis_kelamin"] = gender

        # Nama: halaman 2 dulu, kalau tidak ketemu cari semua halaman
        nama = ""
        if pdf.count >= 2:
            nama = find_field(pdf.text(1), NAME_LABELS)
        if not nama:
            for i in range(pdf.count):
                nama = find_field(pdf.text(i), NAME_LABELS)
                if nama:
                    break
        data["nama_peserta"] = nama

        # No. Telepon: halaman "SURAT ELEGIBILITAS PESERTA" (halaman 2)
        exclude = set()
        for k in ("no_peserta", "no_SEP", "no_rekammedis"):
            v = re.sub(r"\D", "", data.get(k, "") or "")
            if v:
                exclude.add(v)
        data["no_telepon"] = extract_phone(pdf, exclude)

        data["control_ulang"] = ""
        data["dokter"] = ""
        data["ceklis"] = 0

        has_key = bool(data["nama_peserta"] or data["no_peserta"] or data["no_SEP"])
        existing = db.find_existing(data) if has_key else None

        def need(key):
            return existing is None or not str(row_get(existing, key)).strip()

        if has_key:
            if need("control_ulang"):
                data["control_ulang"] = extract_control(pdf)
            # Dokter selalu dibaca dari teks halaman 2 (cepat), supaya data lama yang
            # salah bisa dikoreksi. OCR hanya dipakai kalau kolom dokter masih kosong.
            data["dokter"] = extract_doctor(pdf, allow_ocr=need("dokter"))
        data["_ocr"] = pdf.ocr_used
        return data, existing
    finally:
        pdf.close()


def row_get(row, key):
    """Ambil nilai dari sqlite3.Row dengan aman (kolom tidak ada -> kosong)."""
    try:
        if key in row.keys():
            val = row[key]
            return "" if val is None else val
    except Exception:
        pass
    return ""


def is_done(row):
    try:
        return int(row_get(row, "ceklis") or 0) == 1
    except (ValueError, TypeError):
        return False


def display_value(row, key):
    if key == "ceklis":
        return "SUDAH" if is_done(row) else "BELUM"
    return str(row_get(row, key))


def parse_date(text):
    """Ubah '29/08/2026' atau '29 Agustus 2026' jadi (tahun, bulan, hari). Gagal -> None."""
    text = (text or "").strip()
    m = re.match(r"(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})", text)
    if m:
        d, mo, y = m.groups()
        return (int(y), int(mo), int(d))
    m = re.match(r"(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})", text)
    if m:
        d, mon, y = m.groups()
        mo = month_number(mon)
        if mo:
            return (int(y), mo, int(d))
    return None


def sort_key(col_key, value):
    """Kunci pengurutan: angka sebagai angka, tanggal sebagai tanggal, sisanya teks."""
    text = str(value).strip()
    if col_key in NUMERIC_COLS:
        try:
            return (0, float(text))
        except ValueError:
            return (1, text.lower())
    if col_key in DATE_COLS:
        dt = parse_date(text)
        if dt:
            return (0, dt)
        return (1, text.lower())
    return (0, text.lower()) if text else (1, "")


def default_for(key):
    return 0 if key == "ceklis" else ""


def norm_name(text):
    """Nama untuk perbandingan: huruf besar, spasi ganda dirapikan."""
    return re.sub(r"\s+", " ", str(text or "")).strip().upper()


# ------------------------------------------------------------------- Database
def create_sql(table):
    cols = ",\n".join(f"{k} {t}" for k, _, t in FIELDS)
    return f"CREATE TABLE IF NOT EXISTS {table} (\n id INTEGER PRIMARY KEY AUTOINCREMENT,\n{cols}\n)"


class Database:
    def __init__(self, path):
        self.path = path
        self.conn = sqlite3.connect(path, timeout=30)
        self.conn.row_factory = sqlite3.Row
        self.create()

    def create(self):
        self.conn.execute(create_sql("pasien"))
        self.conn.commit()
        self.migrate()

    def table_cols(self):
        return [r[1] for r in self.conn.execute("PRAGMA table_info(pasien)")]

    def migrate(self):
        """Samakan tabel lama dengan struktur terbaru."""
        cols = self.table_cols()
        lower = {c.lower() for c in cols}

        if "nama_pasien" in lower and "nama_peserta" not in lower:
            self.conn.execute("ALTER TABLE pasien RENAME COLUMN nama_pasien TO nama_peserta")
            self.conn.commit()
            lower = {c.lower() for c in self.table_cols()}

        # hapus kolom umur & tanggal_lahir (tabel dibangun ulang, data lain tetap)
        if lower & OBSOLETE_COLS:
            self.rebuild(lower)
            lower = {c.lower() for c in self.table_cols()}

        # tambah kolom baru (dokter, ceklis, control_ulang) kalau belum ada
        for key, _, typ in FIELDS:
            if key.lower() not in lower:
                self.conn.execute(f"ALTER TABLE pasien ADD COLUMN {key} {typ}")
        self.conn.execute("UPDATE pasien SET ceklis=0 WHERE ceklis IS NULL")
        self.conn.commit()

    def rebuild(self, old_lower):
        """Bangun ulang tabel tanpa kolom lama. File backup dibuat dulu."""
        try:
            bak = self.path + ".bak_v13"
            if os.path.isfile(self.path) and not os.path.exists(bak):
                self.conn.commit()
                shutil.copy2(self.path, bak)
        except Exception:
            pass
        keep = [k for k in ["id"] + DB_FIELDS if k.lower() in old_lower]
        cols = ",".join(keep)
        try:
            self.conn.execute("DROP TABLE IF EXISTS pasien_old")
            self.conn.execute("BEGIN")
            self.conn.execute("ALTER TABLE pasien RENAME TO pasien_old")
            self.conn.execute(create_sql("pasien"))
            self.conn.execute(f"INSERT INTO pasien ({cols}) SELECT {cols} FROM pasien_old")
            self.conn.execute("DROP TABLE pasien_old")
            self.conn.commit()
        except Exception:
            self.conn.rollback()
            raise

    def count(self):
        return self.conn.execute("SELECT COUNT(*) FROM pasien").fetchone()[0]

    def find_existing(self, d):
        """Kembalikan baris yang sama (atau None)."""
        if d.get("no_SEP"):
            cur = self.conn.execute("SELECT * FROM pasien WHERE no_SEP=?", (d["no_SEP"],))
        else:
            cur = self.conn.execute(
                "SELECT * FROM pasien WHERE nama_peserta=? AND no_peserta=? AND tgl_masuk=?",
                (d.get("nama_peserta", ""), d.get("no_peserta", ""), d.get("tgl_masuk", "")),
            )
        return cur.fetchone()

    def update_fields(self, pid, fields):
        """Update sebagian kolom saja (hanya kolom yang valid)."""
        keys = [k for k in fields if k in DB_FIELDS]
        if not keys:
            return
        sets = ",".join(f"{k}=?" for k in keys)
        self.conn.execute(
            f"UPDATE pasien SET {sets} WHERE id=?",
            [fields[k] for k in keys] + [pid],
        )
        self.conn.commit()

    def insert(self, d):
        cols = ",".join(DB_FIELDS)
        qs = ",".join("?" * len(DB_FIELDS))
        self.conn.execute(
            f"INSERT INTO pasien ({cols}) VALUES ({qs})",
            [d.get(k, default_for(k)) for k in DB_FIELDS],
        )
        self.conn.commit()

    def fetch(self, keyword=""):
        if keyword:
            like = f"%{keyword}%"
            where = " OR ".join(f"{k} LIKE ?" for k in SEARCH_FIELDS)
            cur = self.conn.execute(
                f"SELECT * FROM pasien WHERE {where} ORDER BY id DESC",
                (like,) * len(SEARCH_FIELDS),
            )
        else:
            cur = self.conn.execute("SELECT * FROM pasien ORDER BY id DESC")
        return cur.fetchall()

    def get(self, pid):
        return self.conn.execute("SELECT * FROM pasien WHERE id=?", (pid,)).fetchone()

    def update(self, pid, d):
        sets = ",".join(f"{k}=?" for k in DB_FIELDS)
        self.conn.execute(
            f"UPDATE pasien SET {sets} WHERE id=?",
            [d.get(k, default_for(k)) for k in DB_FIELDS] + [pid],
        )
        self.conn.commit()

    def delete(self, pid):
        self.conn.execute("DELETE FROM pasien WHERE id=?", (pid,))
        self.conn.commit()

    def delete_many(self, ids):
        """Hapus banyak data sekaligus (dipecah per 500 supaya aman)."""
        ids = [int(i) for i in ids]
        try:
            for k in range(0, len(ids), 500):
                chunk = ids[k:k + 500]
                self.conn.execute(
                    f"DELETE FROM pasien WHERE id IN ({','.join('?' * len(chunk))})", chunk
                )
            self.conn.commit()
        except Exception:
            self.conn.rollback()
            raise

    # ---------------- gabung database
    def backup_to(self, dest_path):
        """Backup database ini ke file lain (memakai fitur backup resmi SQLite)."""
        self.conn.commit()
        dest = sqlite3.connect(dest_path)
        try:
            self.conn.backup(dest)
        finally:
            dest.close()

    def merge_from(self, src_path):
        """
        Gabungkan data dari file .db lain ke database ini. File sumber hanya dibaca.
        Aturan:
        - NO SEP dan NAMA PESERTA wajib ada isinya, kalau tidak -> dilewati
        - ID sumber tidak dipakai (ID baru dibuat otomatis)
        - NO SEP baru -> baris baru
        - NO SEP sama + nama sama -> data sama: tidak dobel, kolom kosong dilengkapi,
          ceklis tetap SUDAH kalau salah satu sudah
        - NO SEP sama + nama beda -> konflik: tidak digabung, dicatat
        Return (stats, daftar_kosong, daftar_konflik)
        """
        stats = {"total": 0, "baru": 0, "lengkap": 0, "sama": 0, "kosong": 0, "konflik": 0}
        kosong_list, konflik_list = [], []

        uri = Path(src_path).resolve().as_uri() + "?mode=ro"
        src = sqlite3.connect(uri, uri=True, timeout=30)
        src.row_factory = sqlite3.Row
        try:
            has_tbl = src.execute(
                "SELECT name FROM sqlite_master WHERE type='table' AND name='pasien'"
            ).fetchone()
            if not has_tbl:
                raise RuntimeError("Tabel 'pasien' tidak ditemukan di file ini.")
            rows = src.execute("SELECT * FROM pasien ORDER BY id").fetchall()
        finally:
            src.close()

        stats["total"] = len(rows)
        try:
            for r in rows:
                sep = str(row_get(r, "no_SEP")).strip()
                nama = str(row_get(r, "nama_peserta") or row_get(r, "nama_pasien")).strip()
                src_id = row_get(r, "id")

                if not sep or not nama:
                    stats["kosong"] += 1
                    kosong_list.append(
                        f"ID {src_id}: NO SEP='{sep or '-'}' | nama='{nama or '-'}'"
                    )
                    continue

                same_sep = self.conn.execute(
                    "SELECT * FROM pasien WHERE no_SEP=?", (sep,)
                ).fetchall()

                if not same_sep:
                    vals = []
                    for k in DB_FIELDS:
                        if k == "no_SEP":
                            vals.append(sep)
                        elif k == "nama_peserta":
                            vals.append(nama)
                        elif k == "ceklis":
                            vals.append(1 if is_done(r) else 0)
                        else:
                            vals.append(str(row_get(r, k)).strip())
                    self.conn.execute(
                        f"INSERT INTO pasien ({','.join(DB_FIELDS)}) "
                        f"VALUES ({','.join('?' * len(DB_FIELDS))})",
                        vals,
                    )
                    stats["baru"] += 1
                    continue

                match = None
                for x in same_sep:
                    if norm_name(row_get(x, "nama_peserta")) == norm_name(nama):
                        match = x
                        break
                if match is None:
                    stats["konflik"] += 1
                    konflik_list.append(
                        f"NO SEP {sep}: sumber '{nama}' <> database ini "
                        f"'{row_get(same_sep[0], 'nama_peserta')}'"
                    )
                    continue

                # data sama: lengkapi kolom yang kosong saja
                upd = {}
                for k in DB_FIELDS:
                    if k in ("no_SEP", "nama_peserta"):
                        continue
                    if k == "ceklis":
                        if is_done(r) and not is_done(match):
                            upd[k] = 1
                        continue
                    new = str(row_get(r, k)).strip()
                    old = str(row_get(match, k)).strip()
                    if new and not old:
                        upd[k] = new
                if upd:
                    sets = ",".join(f"{k}=?" for k in upd)
                    self.conn.execute(
                        f"UPDATE pasien SET {sets} WHERE id=?",
                        list(upd.values()) + [match["id"]],
                    )
                    stats["lengkap"] += 1
                else:
                    stats["sama"] += 1
            self.conn.commit()
        except Exception:
            self.conn.rollback()
            raise
        return stats, kosong_list, konflik_list

    def close(self):
        try:
            self.conn.close()
        except Exception:
            pass


# ------------------------------------------------------------------ Progress
class ProgressBar(wx.Panel):
    """Loading bar warna hijau + persen."""

    def __init__(self, parent):
        super().__init__(parent, size=(-1, 30))
        self.SetMinSize((-1, 30))
        self.SetBackgroundStyle(wx.BG_STYLE_PAINT)
        self.value = 0.0
        self.text = "Siap"
        self.Bind(wx.EVT_PAINT, self.on_paint)
        self.Bind(wx.EVT_SIZE, self.on_size)

    def on_size(self, e):
        self.Refresh()
        e.Skip()

    def set(self, value, text=None):
        self.value = max(0.0, min(100.0, value))
        if text is not None:
            self.text = text
        self.Refresh()
        self.Update()

    def on_paint(self, _):
        dc = wx.AutoBufferedPaintDC(self)
        w, h = self.GetClientSize()
        dc.SetBackground(wx.Brush(wx.Colour(228, 228, 228)))
        dc.Clear()
        dc.SetPen(wx.TRANSPARENT_PEN)
        dc.SetBrush(wx.Brush(wx.Colour(46, 184, 71)))
        dc.DrawRectangle(0, 0, int(w * self.value / 100.0), h)
        font = self.GetFont()
        font.SetWeight(wx.FONTWEIGHT_BOLD)
        dc.SetFont(font)
        dc.SetTextForeground(wx.Colour(0, 0, 0))
        label = f"{self.text}  {self.value:.0f}%"
        tw, th = dc.GetTextExtent(label)
        dc.DrawText(label, (w - tw) // 2, (h - th) // 2)


# --------------------------------------------------------------- Dialog Ubah
class EditDialog(wx.Dialog):
    def __init__(self, parent, row):
        super().__init__(parent, title="Ubah Data Pasien", style=wx.DEFAULT_DIALOG_STYLE)
        self.ctrls = {}
        grid = wx.FlexGridSizer(0, 2, 8, 10)
        grid.AddGrowableCol(1)
        for key, label, _ in FIELDS:
            val = row_get(row, key)
            grid.Add(wx.StaticText(self, label=label.title()), 0, wx.ALIGN_CENTER_VERTICAL)
            if key == "jenis_kelamin":
                ctrl = wx.Choice(self, choices=GENDERS)
                if val in GENDERS:
                    ctrl.SetSelection(GENDERS.index(val))
            elif key == "ceklis":
                ctrl = wx.CheckBox(self, label="Sudah selesai")
                ctrl.SetValue(is_done(row))
            else:
                ctrl = wx.TextCtrl(self, value=str(val), size=(320, -1))
            self.ctrls[key] = ctrl
            grid.Add(ctrl, 1, wx.EXPAND)

        main = wx.BoxSizer(wx.VERTICAL)
        main.Add(grid, 1, wx.ALL | wx.EXPAND, 15)
        main.Add(self.CreateStdDialogButtonSizer(wx.OK | wx.CANCEL), 0, wx.ALL | wx.ALIGN_RIGHT, 10)
        self.SetSizerAndFit(main)
        self.CentreOnParent()

    def values(self):
        out = {}
        for key, ctrl in self.ctrls.items():
            if isinstance(ctrl, wx.Choice):
                idx = ctrl.GetSelection()
                out[key] = GENDERS[idx] if idx != wx.NOT_FOUND else ""
            elif isinstance(ctrl, wx.CheckBox):
                out[key] = 1 if ctrl.GetValue() else 0
            else:
                out[key] = ctrl.GetValue().strip()
        return out


# ------------------------------------------------- Dialog pilih banyak folder
class MultiFolderDialog(wx.Dialog):
    """Pilih banyak folder sekaligus: Ctrl+klik (terpisah) atau Shift+klik (blok berurutan)."""

    def __init__(self, parent, start_dir=""):
        super().__init__(
            parent, title="Pilih Folder PDF (boleh banyak)", size=(640, 620),
            style=wx.DEFAULT_DIALOG_STYLE | wx.RESIZE_BORDER,
        )
        self.extra = []  # folder tambahan dari tombol "+ Folder lain..."

        style = wx.DIRCTRL_DIR_ONLY | getattr(wx, "DIRCTRL_MULTIPLE", 0)
        kw = {}
        if start_dir and os.path.isdir(start_dir):
            kw["dir"] = start_dir
        self.tree = wx.GenericDirCtrl(self, style=style, **kw)

        info = wx.StaticText(
            self,
            label=(
                "Ctrl + klik   : pilih beberapa folder terpisah\n"
                "Shift + klik  : pilih blok folder berurutan\n"
                "Subfolder di dalam folder terpilih otomatis ikut di-scan."
            ),
        )
        self.lbl_extra = wx.StaticText(self, label="")
        btn_extra = wx.Button(self, label="+ Folder lain...")
        btn_extra.Bind(wx.EVT_BUTTON, self.on_add_extra)

        row = wx.BoxSizer(wx.HORIZONTAL)
        row.Add(btn_extra, 0, wx.RIGHT, 10)
        row.Add(self.lbl_extra, 1, wx.ALIGN_CENTER_VERTICAL)

        main = wx.BoxSizer(wx.VERTICAL)
        main.Add(info, 0, wx.ALL, 10)
        main.Add(self.tree, 1, wx.LEFT | wx.RIGHT | wx.EXPAND, 10)
        main.Add(row, 0, wx.ALL | wx.EXPAND, 10)
        main.Add(self.CreateStdDialogButtonSizer(wx.OK | wx.CANCEL), 0, wx.ALL | wx.ALIGN_RIGHT, 10)
        self.SetSizer(main)
        self.Bind(wx.EVT_BUTTON, self.on_ok, id=wx.ID_OK)
        self.CentreOnParent()

    def on_add_extra(self, _):
        dlg = wx.DirDialog(self, "Pilih folder tambahan", style=wx.DD_DEFAULT_STYLE)
        if dlg.ShowModal() == wx.ID_OK:
            p = dlg.GetPath()
            if p and p not in self.extra:
                self.extra.append(p)
            self.lbl_extra.SetLabel(f"Folder tambahan: {len(self.extra)}")
            self.Layout()
        dlg.Destroy()

    def get_paths(self):
        paths = []
        try:
            paths = list(self.tree.GetPaths())
        except Exception:
            paths = []
        if not paths:
            try:
                tc = self.tree.GetTreeCtrl()
                for it in tc.GetSelections():
                    paths.append(self.tree.GetPath(it))
            except Exception:
                pass
        if not paths:
            try:
                p = self.tree.GetPath()
                if p:
                    paths = [p]
            except Exception:
                pass
        paths += self.extra
        return [p for p in dict.fromkeys(paths) if p and os.path.isdir(p)]

    def on_ok(self, e):
        if not self.get_paths():
            wx.MessageBox("Pilih minimal satu folder dulu.", "Pilih Folder",
                          wx.OK | wx.ICON_INFORMATION)
            return
        e.Skip()


# ---------------------------------------------------------------- Main Frame
class MainFrame(wx.Frame):
    def __init__(self):
        super().__init__(None, title=APP_TITLE, size=(1350, 720))
        self.stop = threading.Event()  # tombol BATAL PROSES -> stop.set()
        self.worker = None
        self.sort_col = None   # index kolom yang sedang di-sort
        self.sort_asc = True
        self.search_timer = None
        self.current_rows = []  # baris yang sedang tampil di list (dipakai export)

        # buka database terakhir yang dipilih (atau default di folder .exe)
        self.db_path = load_last_db()
        try:
            self.db = Database(self.db_path)
        except Exception as ex:
            wx.MessageBox(
                f"Gagal membuka database:\n{self.db_path}\n\n{ex}\n\nMemakai database default.",
                "Error", wx.OK | wx.ICON_ERROR,
            )
            self.db_path = DEFAULT_DB
            self.db = Database(self.db_path)
        self.SetTitle(f"{APP_TITLE} - {self.db_path}")

        main = wx.BoxSizer(wx.HORIZONTAL)
        self.left = self.build_left()
        self.right = self.build_right()
        main.Add(self.left, 2, wx.EXPAND)   # 20%
        main.Add(self.right, 8, wx.EXPAND)  # 80%
        self.SetSizer(main)

        self.Bind(wx.EVT_CLOSE, self.on_close)
        self.Centre()
        self.Maximize(True)  # otomatis maximize, tetap bisa minimize / restore down
        self.refresh_grid()

    # ---------------- sisi kiri (20%)
    def build_left(self):
        p = wx.Panel(self)
        p.SetBackgroundColour(wx.Colour(245, 247, 250))
        s = wx.BoxSizer(wx.VERTICAL)

        title = wx.StaticText(p, label="SELEKSI FOLDER PDF")
        f = title.GetFont()
        f.SetWeight(wx.FONTWEIGHT_BOLD)
        f.SetPointSize(f.GetPointSize() + 2)
        title.SetFont(f)
        s.Add(title, 0, wx.ALL, 12)

        self.btn_folder = wx.Button(p, label="PILIH FOLDER (+ SUBFOLDER)", size=(-1, 40))
        self.btn_folder.Bind(wx.EVT_BUTTON, self.on_pick_folder)
        s.Add(self.btn_folder, 0, wx.LEFT | wx.RIGHT | wx.EXPAND, 12)

        self.btn_cancel = wx.Button(p, label="BATALKAN PROSES")
        self.btn_cancel.Disable()
        self.btn_cancel.Bind(wx.EVT_BUTTON, self.on_cancel)
        s.Add(self.btn_cancel, 0, wx.LEFT | wx.RIGHT | wx.TOP | wx.EXPAND, 12)

        self.lbl_folder = wx.StaticText(p, label="Folder: -", style=wx.ST_ELLIPSIZE_MIDDLE)
        s.Add(self.lbl_folder, 0, wx.ALL | wx.EXPAND, 12)

        self.lbl_stat = wx.StaticText(p, label=self.stat_text(0, 0, 0, 0))
        s.Add(self.lbl_stat, 0, wx.LEFT | wx.RIGHT | wx.BOTTOM, 12)

        self.log = wx.TextCtrl(p, style=wx.TE_MULTILINE | wx.TE_READONLY | wx.HSCROLL)
        s.Add(self.log, 1, wx.LEFT | wx.RIGHT | wx.BOTTOM | wx.EXPAND, 12)

        p.SetSizer(s)
        return p

    @staticmethod
    def stat_text(baru, upd, dup, gagal):
        return f"Baru: {baru}   Update: {upd}\nDuplikat: {dup}   Gagal: {gagal}"

    # ---------------- sisi kanan (80%)
    def build_right(self):
        p = wx.Panel(self)
        s = wx.BoxSizer(wx.VERTICAL)

        bar = wx.BoxSizer(wx.HORIZONTAL)
        self.btn_db = wx.Button(p, label="SELEKSI DATABASE")
        self.btn_merge = wx.Button(p, label="GABUNG DATABASE")
        self.btn_merge.SetBackgroundColour(wx.Colour(0, 102, 204))
        self.btn_merge.SetForegroundColour(wx.Colour(255, 255, 255))
        self.btn_edit = wx.Button(p, label="UBAH DATA")
        self.btn_del = wx.Button(p, label="HAPUS")
        self.btn_ref = wx.Button(p, label="REFRESH")
        self.btn_export = wx.Button(p, label="EXPORT TO EXCEL")
        self.btn_export.SetBackgroundColour(wx.Colour(33, 115, 70))
        self.btn_export.SetForegroundColour(wx.Colour(255, 255, 255))
        self.txt_search = wx.TextCtrl(p, size=(300, -1))
        self.txt_search.SetHint("Cari: nama / no. peserta / SEP / telp / dokter...")
        self.lbl_total = wx.StaticText(p, label="Total: 0")

        for b in (self.btn_db, self.btn_merge, self.btn_edit, self.btn_del, self.btn_ref, self.btn_export):
            bar.Add(b, 0, wx.RIGHT, 6)
        bar.AddStretchSpacer(1)
        bar.Add(self.lbl_total, 0, wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 10)
        bar.Add(self.txt_search, 0)
        s.Add(bar, 0, wx.ALL | wx.EXPAND, 8)

        self.grid = gridlib.Grid(p)
        self.grid.CreateGrid(0, len(GRID_COLS))
        for i, (key, _) in enumerate(GRID_COLS):
            self.grid.SetColSize(i, COL_WIDTH.get(key, 120))
        self.grid.SetRowLabelSize(0)
        self.grid.EnableEditing(False)
        # mode pilih baris; Ctrl+klik / Shift+klik / drag = pilih banyak baris
        self.grid.SetSelectionMode(getattr(gridlib.Grid, "GridSelectRows", getattr(gridlib.Grid, "SelectRows", 1)))

        # ---- tinggi baris dikunci: kursor tidak berubah jadi panah atas-bawah,
        #      baris tidak bisa ditarik / di-resize (lebar kolom tetap bisa)
        self.grid.SetDefaultRowSize(ROW_HEIGHT, True)
        try:
            self.grid.DisableDragRowSize()
        except Exception:
            pass
        try:
            self.grid.EnableDragRowSize(False)
        except Exception:
            pass
        try:
            self.grid.DisableDragGridSize()
        except Exception:
            pass
        try:
            self.grid.EnableDragColSize(True)  # lebar kolom tetap boleh diatur
        except Exception:
            pass

        self.update_header_labels()

        # kolom WHATSAPP: hijau, putih, tebal, rata tengah
        attr = gridlib.GridCellAttr()
        attr.SetBackgroundColour(COLOR_WA)
        attr.SetTextColour(wx.Colour(255, 255, 255))
        attr.SetAlignment(wx.ALIGN_CENTRE, wx.ALIGN_CENTRE)
        fnt = wx.Font(self.grid.GetDefaultCellFont())
        fnt.SetWeight(wx.FONTWEIGHT_BOLD)
        attr.SetFont(fnt)
        self.grid.SetColAttr(WA_COL, attr)

        # kolom CEKLIS: tebal, rata tengah (warna diatur per baris)
        attr2 = gridlib.GridCellAttr()
        attr2.SetAlignment(wx.ALIGN_CENTRE, wx.ALIGN_CENTRE)
        fnt2 = wx.Font(self.grid.GetDefaultCellFont())
        fnt2.SetWeight(wx.FONTWEIGHT_BOLD)
        attr2.SetFont(fnt2)
        self.grid.SetColAttr(CEK_COL, attr2)

        s.Add(self.grid, 1, wx.LEFT | wx.RIGHT | wx.EXPAND, 8)

        self.progress = ProgressBar(p)
        s.Add(self.progress, 0, wx.ALL | wx.EXPAND, 8)

        self.btn_db.Bind(wx.EVT_BUTTON, self.on_select_db)
        self.btn_merge.Bind(wx.EVT_BUTTON, self.on_merge_db)
        self.btn_edit.Bind(wx.EVT_BUTTON, self.on_edit)
        self.btn_del.Bind(wx.EVT_BUTTON, self.on_delete)
        self.btn_ref.Bind(wx.EVT_BUTTON, lambda e: self.refresh_grid())
        self.btn_export.Bind(wx.EVT_BUTTON, self.on_export_excel)
        self.txt_search.Bind(wx.EVT_TEXT, self.on_search_text)
        self.grid.Bind(gridlib.EVT_GRID_CELL_LEFT_CLICK, self.on_cell_click)
        self.grid.Bind(gridlib.EVT_GRID_CELL_LEFT_DCLICK, self.on_cell_dclick)
        self.grid.Bind(gridlib.EVT_GRID_LABEL_LEFT_CLICK, self.on_header_click)
        self.grid.Bind(gridlib.EVT_GRID_ROW_SIZE, lambda e: e.Veto())  # cadangan: tolak resize baris
        self.grid.Bind(wx.EVT_KEY_DOWN, self.on_grid_key_down)  # DELETE, Backspace, Esc
        self.grid.Bind(wx.EVT_CHAR, self.on_grid_char)          # ketik langsung -> kolom pencarian

        p.SetSizer(s)
        return p

    # ---------------- keyboard di tabel (v17)
    def on_grid_key_down(self, e):
        code = e.GetKeyCode()
        if code in (wx.WXK_DELETE, wx.WXK_NUMPAD_DELETE):
            self.delete_selected(use_cursor=False)
            return
        if code == wx.WXK_BACK and not (e.ControlDown() or e.AltDown()):
            cur = self.txt_search.GetValue()
            if cur:
                self.txt_search.SetValue(cur[:-1])
                self.txt_search.SetInsertionPointEnd()
                return
        if code == wx.WXK_ESCAPE and self.txt_search.GetValue():
            self.txt_search.SetValue("")
            return
        e.Skip()

    def on_grid_char(self, e):
        # Ctrl+huruf / Alt+huruf dibiarkan (AltGr = Ctrl+Alt tetap dianggap mengetik)
        if e.ControlDown() != e.AltDown():
            e.Skip()
            return
        ch = e.GetUnicodeKey()
        if ch and ch >= 32 and ch != 127:
            self.txt_search.SetFocus()
            self.txt_search.AppendText(chr(ch))
            self.txt_search.SetInsertionPointEnd()
            return
        e.Skip()

    # ---------------- pencarian langsung
    def on_search_text(self, _):
        if self.search_timer is not None:
            self.search_timer.Stop()
        self.search_timer = wx.CallLater(150, self.refresh_grid)

    # ---------------- sort header
    def update_header_labels(self):
        for i, (_, label) in enumerate(GRID_COLS):
            arrow = ""
            if i == self.sort_col:
                arrow = "  ▲" if self.sort_asc else "  ▼"
            self.grid.SetColLabelValue(i, label + arrow)

    def on_header_click(self, e):
        col = e.GetCol()
        if col < 0 or col >= len(GRID_COLS) or GRID_COLS[col][0] == "wa":
            return
        if self.sort_col == col:
            self.sort_asc = not self.sort_asc
        else:
            self.sort_col = col
            self.sort_asc = True
        self.update_header_labels()
        self.refresh_grid()

    def sort_rows(self, rows):
        if self.sort_col is None:
            return list(rows)
        key = GRID_COLS[self.sort_col][0]
        return sorted(rows, key=lambda r: sort_key(key, row_get(r, key)), reverse=not self.sort_asc)

    # ---------------- grid
    def paint_row(self, r, row):
        """Isi & warnai satu baris. Baris yang sudah di-ceklis -> hijau tipis."""
        g = self.grid
        done = is_done(row)
        bg = COLOR_DONE if done else (COLOR_WHITE if r % 2 == 0 else COLOR_GREY)
        for c, (key, _) in enumerate(GRID_COLS):
            if key == "wa":
                g.SetCellValue(r, c, "WhatsApp")
                g.SetCellBackgroundColour(r, c, COLOR_WA)
                g.SetCellTextColour(r, c, COLOR_WHITE)
                continue
            if key == "ceklis":
                g.SetCellValue(r, c, ("✔ " if done else "☐ ") + display_value(row, key))
                g.SetCellTextColour(r, c, COLOR_DONE_TEXT if done else COLOR_BLACK)
            else:
                g.SetCellValue(r, c, display_value(row, key))
            g.SetCellBackgroundColour(r, c, bg)

    def refresh_grid(self):
        try:
            rows = self.sort_rows(self.db.fetch(self.txt_search.GetValue().strip()))
        except Exception as ex:
            wx.MessageBox(f"Gagal membaca database:\n{ex}", "Error", wx.OK | wx.ICON_ERROR)
            return
        self.current_rows = rows  # simpan baris yang tampil, dipakai untuk export
        g = self.grid
        g.BeginBatch()
        n = g.GetNumberRows()
        if n:
            g.DeleteRows(0, n)
        g.AppendRows(len(rows))
        for r, row in enumerate(rows):
            self.paint_row(r, row)
        g.EndBatch()
        self.lbl_total.SetLabel(f"Total: {len(rows)}")
        self.right.Layout()

    def selected_rows(self, use_cursor=True):
        """Nomor baris (index tabel) yang sedang diseleksi, mendukung multi seleksi."""
        g = self.grid
        rows = set(g.GetSelectedRows())
        try:
            tls = g.GetSelectionBlockTopLeft()
            brs = g.GetSelectionBlockBottomRight()
            for tl, br in zip(tls, brs):
                rows.update(range(tl.GetRow(), br.GetRow() + 1))
        except Exception:
            pass
        if not rows and use_cursor:
            r = g.GetGridCursorRow()
            if r is not None and 0 <= r < g.GetNumberRows():
                rows.add(r)
        return sorted(r for r in rows if 0 <= r < g.GetNumberRows())

    def selected_ids(self, use_cursor=True):
        ids = []
        for r in self.selected_rows(use_cursor):
            try:
                ids.append(int(self.grid.GetCellValue(r, 0)))
            except ValueError:
                pass
        return ids

    def selected_id(self):
        ids = self.selected_ids(True)
        return ids[0] if ids else None

    def on_cell_click(self, e):
        # Ctrl / Shift + klik = memilih banyak baris, jangan jalankan aksi tombol
        if e.ShiftDown() or e.ControlDown():
            e.Skip()
            return
        row, col = e.GetRow(), e.GetCol()
        if 0 <= col < len(GRID_COLS):
            key = GRID_COLS[col][0]
            if key == "wa":
                self.open_whatsapp(self.grid.GetCellValue(row, PHONE_COL))
            elif key == "ceklis":
                self.toggle_ceklis(row)
        e.Skip()

    def toggle_ceklis(self, r):
        """Klik CEKLIS: belum <-> sudah, disimpan ke database, baris jadi hijau tipis."""
        try:
            pid = int(self.grid.GetCellValue(r, 0))
        except ValueError:
            return
        row = self.db.get(pid)
        if not row:
            return
        new_val = 0 if is_done(row) else 1
        self.db.update_fields(pid, {"ceklis": new_val})
        fresh = self.db.get(pid)
        if r < len(self.current_rows) and self.current_rows[r]["id"] == pid:
            self.current_rows[r] = fresh
        self.paint_row(r, fresh)
        self.grid.ForceRefresh()

    def on_cell_dclick(self, e):
        """Double klik baris -> UBAH DATA (kolom WhatsApp & Ceklis diabaikan)."""
        row, col = e.GetRow(), e.GetCol()
        if col in (WA_COL, CEK_COL):
            return
        try:
            pid = int(self.grid.GetCellValue(row, 0))
        except ValueError:
            return
        self.grid.SelectRow(row)
        self.edit_by_id(pid)

    def open_whatsapp(self, phone):
        digits = re.sub(r"\D", "", phone or "")
        if not digits:
            wx.MessageBox("Nomor telepon kosong.", "WhatsApp", wx.OK | wx.ICON_WARNING)
            return
        if digits.startswith("0"):
            digits = "62" + digits[1:]
        elif digits.startswith("8"):
            digits = "62" + digits
        webbrowser.open(f"https://wa.me/{digits}")

    # ---------------- export excel
    def on_export_excel(self, _):
        """Export data yang sedang tampil di list (sesuai pencarian & urutan)."""
        if openpyxl is None:
            wx.MessageBox(
                "Library Excel belum terpasang.\n\nJalankan di CMD:\npip install openpyxl",
                "Export to Excel", wx.OK | wx.ICON_WARNING,
            )
            return
        rows = self.current_rows
        if not rows:
            wx.MessageBox("Tidak ada data untuk diexport.", "Export to Excel", wx.OK | wx.ICON_INFORMATION)
            return

        keyword = self.txt_search.GetValue().strip()
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        default_name = f"Data_Pasien_{'Hasil_Pencarian_' if keyword else ''}{stamp}.xlsx"

        dlg = wx.FileDialog(
            self, "Simpan sebagai Excel", defaultDir=BASE_DIR, defaultFile=default_name,
            wildcard="Excel (*.xlsx)|*.xlsx", style=wx.FD_SAVE | wx.FD_OVERWRITE_PROMPT,
        )
        if dlg.ShowModal() != wx.ID_OK:
            dlg.Destroy()
            return
        path = dlg.GetPath()
        dlg.Destroy()
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Data Pasien"

            # header
            ws.append([LABELS[k] for k in EXPORT_KEYS])
            head_fill = PatternFill("solid", fgColor="2EB847")
            for c in range(1, len(EXPORT_KEYS) + 1):
                cell = ws.cell(row=1, column=c)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = head_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # isi data (semua disimpan sebagai teks supaya angka 0 di depan tidak hilang)
            done_fill = PatternFill("solid", fgColor="C6EFCE")
            for row in rows:
                ws.append([display_value(row, k) for k in EXPORT_KEYS])
            for r, row in enumerate(rows, 2):
                for c in range(1, len(EXPORT_KEYS) + 1):
                    cell = ws.cell(row=r, column=c)
                    cell.number_format = "@"
                    if is_done(row):
                        cell.fill = done_fill  # baris sudah selesai -> hijau tipis

            # lebar kolom otomatis
            for c, key in enumerate(EXPORT_KEYS, 1):
                longest = len(LABELS[key])
                for row in rows:
                    longest = max(longest, len(display_value(row, key)))
                ws.column_dimensions[get_column_letter(c)].width = min(longest + 3, 45)

            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            wb.save(path)
        except PermissionError:
            wx.MessageBox(
                "File tidak bisa disimpan. Kemungkinan file itu sedang terbuka di Excel.\n"
                "Tutup dulu file-nya atau pakai nama lain.",
                "Export to Excel", wx.OK | wx.ICON_ERROR,
            )
            return
        except Exception as ex:
            wx.MessageBox(f"Gagal export:\n{ex}", "Export to Excel", wx.OK | wx.ICON_ERROR)
            return

        self.progress.set(100, f"Export selesai: {len(rows)} data")
        ask = wx.MessageBox(
            f"Berhasil export {len(rows)} data ke:\n{path}\n\nBuka file sekarang?",
            "Export to Excel", wx.YES_NO | wx.ICON_INFORMATION,
        )
        if ask == wx.YES:
            try:
                os.startfile(path)  # Windows
            except Exception:
                pass

    # ---------------- tombol
    def on_select_db(self, _):
        start_dir = os.path.dirname(self.db_path) if os.path.isdir(os.path.dirname(self.db_path)) else BASE_DIR
        dlg = wx.FileDialog(
            self, "Pilih / buat database", defaultDir=start_dir,
            defaultFile=os.path.basename(self.db_path),
            wildcard="SQLite DB (*.db)|*.db|Semua file (*.*)|*.*", style=wx.FD_SAVE,
        )
        if dlg.ShowModal() == wx.ID_OK:
            path = dlg.GetPath()
            if not os.path.exists(path) and not os.path.splitext(path)[1]:
                path += ".db"
            try:
                new_db = Database(path)
                total = new_db.count()  # tes baca tabel
            except Exception as ex:
                wx.MessageBox(f"Gagal membuka database:\n{path}\n\n{ex}", "Error", wx.OK | wx.ICON_ERROR)
                dlg.Destroy()
                return
            self.db.close()
            self.db = new_db
            self.db_path = path
            save_last_db(path)  # diingat untuk pembukaan berikutnya
            self.SetTitle(f"{APP_TITLE} - {path}")
            self.refresh_grid()
            self.progress.set(0, f"Database dibuka: {total} data")
        dlg.Destroy()

    # ---------------- gabung database
    def on_merge_db(self, _):
        """Gabungkan satu atau banyak file .db ke database yang sedang dibuka."""
        if self.worker and self.worker.is_alive():
            wx.MessageBox("Tunggu proses scan folder selesai dulu.", "Gabung Database",
                          wx.OK | wx.ICON_INFORMATION)
            return

        dlg = wx.FileDialog(
            self, "Pilih file database yang mau digabung (boleh banyak)",
            defaultDir=BASE_DIR,
            wildcard="SQLite DB (*.db)|*.db|Semua file (*.*)|*.*",
            style=wx.FD_OPEN | wx.FD_FILE_MUST_EXIST | wx.FD_MULTIPLE,
        )
        if dlg.ShowModal() != wx.ID_OK:
            dlg.Destroy()
            return
        paths = list(dlg.GetPaths())
        dlg.Destroy()

        master = os.path.normcase(os.path.abspath(self.db_path))
        sources = []
        for p in paths:
            if os.path.normcase(os.path.abspath(p)) == master:
                self.log.AppendText(f"[LEWAT] {os.path.basename(p)} - itu database yang sedang dibuka\n")
            else:
                sources.append(p)
        if not sources:
            wx.MessageBox("Tidak ada file yang bisa digabung.", "Gabung Database",
                          wx.OK | wx.ICON_INFORMATION)
            return

        names = "\n".join("  - " + os.path.basename(p) for p in sources)
        ask = wx.MessageBox(
            f"Gabungkan {len(sources)} file ini ke:\n{self.db_path}\n\n{names}\n\n"
            "Database ini akan di-backup otomatis dulu. File sumber tidak diubah.\n"
            "Lanjut?",
            "Gabung Database", wx.YES_NO | wx.ICON_QUESTION,
        )
        if ask != wx.YES:
            return

        # backup otomatis sebelum gabung
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        bak = f"{self.db_path}.bak_gabung_{stamp}"
        try:
            self.db.backup_to(bak)
        except Exception as ex:
            wx.MessageBox(f"Backup gagal, proses gabung dibatalkan:\n{ex}", "Gabung Database",
                          wx.OK | wx.ICON_ERROR)
            return

        self.log.Clear()
        self.log.AppendText(f"=== GABUNG DATABASE ===\nBackup: {os.path.basename(bak)}\n\n")

        tot = {"total": 0, "baru": 0, "lengkap": 0, "sama": 0, "kosong": 0, "konflik": 0}
        file_gagal = 0
        wx.BeginBusyCursor()
        try:
            for n, p in enumerate(sources, 1):
                name = os.path.basename(p)
                self.progress.set((n - 1) * 100.0 / len(sources), f"Gabung {n}/{len(sources)}")
                try:
                    st, kosong_list, konflik_list = self.db.merge_from(p)
                except Exception as ex:
                    file_gagal += 1
                    self.log.AppendText(f"[GAGAL] {name} - {ex}\n\n")
                    continue
                for k in tot:
                    tot[k] += st[k]
                self.log.AppendText(
                    f"[OK] {name}\n"
                    f"     isi file : {st['total']} data\n"
                    f"     baru     : {st['baru']}\n"
                    f"     dilengkapi: {st['lengkap']}\n"
                    f"     sama     : {st['sama']}\n"
                    f"     dilewati (NO SEP / nama kosong): {st['kosong']}\n"
                    f"     konflik (NO SEP sama, nama beda): {st['konflik']}\n"
                )
                for line in kosong_list[:MAX_LOG_LIST]:
                    self.log.AppendText(f"       [KOSONG ] {line}\n")
                if len(kosong_list) > MAX_LOG_LIST:
                    self.log.AppendText(f"       ... dan {len(kosong_list) - MAX_LOG_LIST} lainnya\n")
                for line in konflik_list[:MAX_LOG_LIST]:
                    self.log.AppendText(f"       [KONFLIK] {line}\n")
                if len(konflik_list) > MAX_LOG_LIST:
                    self.log.AppendText(f"       ... dan {len(konflik_list) - MAX_LOG_LIST} lainnya\n")
                self.log.AppendText("\n")
        finally:
            wx.EndBusyCursor()

        self.progress.set(100, "Gabung selesai")
        self.lbl_stat.SetLabel(self.stat_text(tot["baru"], tot["lengkap"], tot["sama"], tot["konflik"] + tot["kosong"]))
        self.log.AppendText(
            "--- RINGKASAN ---\n"
            f"Data baru masuk : {tot['baru']}\n"
            f"Dilengkapi      : {tot['lengkap']}\n"
            f"Sama (dilewati) : {tot['sama']}\n"
            f"NO SEP/nama kosong: {tot['kosong']}\n"
            f"Konflik         : {tot['konflik']}\n"
            f"File gagal      : {file_gagal}\n"
        )
        self.refresh_grid()
        wx.MessageBox(
            f"Gabung selesai.\n\n"
            f"Data baru masuk : {tot['baru']}\n"
            f"Dilengkapi : {tot['lengkap']}\n"
            f"Sama (tidak dobel) : {tot['sama']}\n"
            f"Dilewati (NO SEP / nama kosong) : {tot['kosong']}\n"
            f"Konflik (NO SEP sama, nama beda) : {tot['konflik']}\n"
            f"File gagal : {file_gagal}\n\n"
            "Detail ada di kotak log sebelah kiri.",
            "Gabung Database", wx.OK | wx.ICON_INFORMATION,
        )

    def edit_by_id(self, pid):
        row = self.db.get(pid)
        if not row:
            return
        dlg = EditDialog(self, row)
        if dlg.ShowModal() == wx.ID_OK:
            self.db.update(pid, dlg.values())
            self.refresh_grid()
        dlg.Destroy()

    def on_edit(self, _):
        pid = self.selected_id()
        if pid is None:
            wx.MessageBox("Klik salah satu data di list dulu.", "Ubah Data", wx.OK | wx.ICON_INFORMATION)
            return
        self.edit_by_id(pid)

    # ---------------- hapus (satu / banyak data)
    def delete_selected(self, use_cursor=True):
        ids = self.selected_ids(use_cursor)
        if not ids:
            wx.MessageBox("Klik / seleksi data di list dulu.", "Hapus", wx.OK | wx.ICON_INFORMATION)
            return
        if len(ids) == 1:
            row = self.db.get(ids[0])
            nama = row_get(row, "nama_peserta") if row else ""
            sep = row_get(row, "no_SEP") if row else ""
            text = f"Hapus data ini?\n\n{nama}\nNO SEP: {sep}"
        else:
            text = f"Hapus {len(ids)} data yang dipilih?"
        dlg = wx.MessageDialog(self, text, "Konfirmasi Hapus",
                               wx.YES_NO | wx.NO_DEFAULT | wx.ICON_QUESTION)
        try:
            dlg.SetYesNoLabels("Ya", "Tidak")
        except Exception:
            pass
        ok = dlg.ShowModal() == wx.ID_YES
        dlg.Destroy()
        if not ok:
            return
        try:
            self.db.delete_many(ids)
        except Exception as ex:
            wx.MessageBox(f"Gagal menghapus:\n{ex}", "Hapus", wx.OK | wx.ICON_ERROR)
            return
        self.refresh_grid()
        self.progress.set(self.progress.value, f"{len(ids)} data dihapus")

    def on_delete(self, _):
        self.delete_selected(use_cursor=True)

    # ---------------- scan folder
    def post(self, fn, *args):
        """Kirim hasil dari thread kerja ke tampilan (aman kalau jendela sudah ditutup)."""
        try:
            wx.CallAfter(fn, *args)
        except Exception:
            pass

    def on_pick_folder(self, _):
        if self.worker and self.worker.is_alive():
            return
        dlg = MultiFolderDialog(self, load_config().get("last_folder", ""))
        if dlg.ShowModal() != wx.ID_OK:
            dlg.Destroy()
            return
        folders = dlg.get_paths()
        dlg.Destroy()
        if not folders:
            return

        save_config(last_folder=folders[0])
        if len(folders) == 1:
            self.lbl_folder.SetLabel(f"Folder: {folders[0]}")
        else:
            self.lbl_folder.SetLabel(f"{len(folders)} folder dipilih")
        self.lbl_folder.SetToolTip("\n".join(folders))

        self.log.Clear()
        self.log.AppendText("Folder yang di-scan:\n" + "\n".join("  - " + f for f in folders) + "\n\n")
        if not OCR_AVAILABLE:
            self.log.AppendText(
                "[INFO] OCR belum terpasang, halaman hasil scan tidak bisa dibaca.\n"
                "       Jalankan di CMD: pip install rapidocr-onnxruntime\n\n"
            )
        self.stop.clear()
        self.btn_folder.Disable()
        self.btn_merge.Disable()
        self.btn_cancel.Enable()
        self.progress.set(0, "Memulai...")
        self.worker = threading.Thread(target=self.scan_worker, args=(folders,), daemon=True)
        self.worker.start()

    def on_cancel(self, _):
        """BATAL PROSES: langsung bereaksi, thread kerja berhenti di titik cek terdekat."""
        if self.worker and self.worker.is_alive() and not self.stop.is_set():
            self.stop.set()
            self.btn_cancel.Disable()
            self.progress.set(self.progress.value, "Membatalkan...")
            self.log.AppendText("[INFO] Pembatalan diminta, berhenti sebentar lagi...\n")

    def scan_worker(self, folders):
        db = None
        baru = upd = dup = gagal = 0
        processed = 0
        total = 0
        try:
            db = Database(self.db_path)  # koneksi sendiri untuk thread ini

            # kumpulkan semua PDF dari semua folder (file yang sama tidak dobel)
            found = {}
            for folder in folders:
                for root, _, names in os.walk(folder):
                    if self.stop.is_set():
                        break
                    for n in names:
                        if n.lower().endswith(".pdf"):
                            full = os.path.join(root, n)
                            found.setdefault(os.path.normcase(os.path.abspath(full)), full)
                if self.stop.is_set():
                    break
            files = sorted(found.values())
            total = len(files)

            if self.stop.is_set():
                self.post(self.on_done, 0, 0, 0, 0, "Dibatalkan")
                return
            if total == 0:
                self.post(self.on_done, 0, 0, 0, 0, "Tidak ada file PDF.")
                return

            self.post(self.on_progress, 0, total, 0, 0, 0, 0, f"Ditemukan {total} file PDF.", False)

            fill_fields = [("no_telepon", "telp"), ("control_ulang", "kontrol"), ("dokter", "dokter")]
            buf = []
            last_flush = time.time()
            last_refresh = time.time()

            for i, path in enumerate(files, 1):
                if self.stop.is_set():
                    break
                name = os.path.basename(path)
                try:
                    data, existing = extract_patient(path, db, self.stop)
                    if not data or not (data["nama_peserta"] or data["no_peserta"] or data["no_SEP"]):
                        gagal += 1
                        msg = f"[GAGAL] {name} - data tidak ditemukan"
                    else:
                        tag = " [OCR]" if data.get("_ocr") else ""
                        if existing is None:
                            db.insert(data)
                            baru += 1
                            msg = (
                                f"[BARU ] {name} - {data['nama_peserta']} | "
                                f"telp: {data['no_telepon'] or '-'} | "
                                f"kontrol: {data['control_ulang'] or '-'} | "
                                f"dokter: {data['dokter'] or '-'}{tag}"
                            )
                        else:
                            # isi data yang masih kosong di database.
                            # khusus DOKTER: kalau beda dengan isi PDF halaman 2, ditimpa (koreksi data lama)
                            fields = {}
                            notes = []
                            for key, lab in fill_fields:
                                old = str(row_get(existing, key)).strip()
                                new = str(data.get(key) or "").strip()
                                if not new:
                                    continue
                                if key == "dokter":
                                    if new != old:
                                        fields[key] = new
                                        notes.append(f"{lab}: {new}")
                                elif not old:
                                    fields[key] = new
                                    notes.append(f"{lab}: {new}")
                            if fields:
                                db.update_fields(existing["id"], fields)
                                upd += 1
                                msg = f"[UPDATE] {name} - diisi {', '.join(notes)}{tag}"
                            else:
                                dup += 1
                                msg = f"[SKIP ] {name} - sudah ada ({data['nama_peserta']}){tag}"
                except Cancelled:
                    break
                except Exception as ex:
                    gagal += 1
                    msg = f"[ERROR] {name} - {ex}"

                processed = i
                buf.append(msg)
                now = time.time()
                if now - last_flush >= 0.2 or i == total:
                    refresh = (now - last_refresh) >= 3.0
                    if refresh:
                        last_refresh = now
                    self.post(self.on_progress, i, total, baru, upd, dup, gagal, "\n".join(buf), refresh)
                    buf = []
                    last_flush = now

            if buf:  # sisa log yang belum terkirim (misal karena dibatalkan)
                self.post(self.on_progress, processed, total, baru, upd, dup, gagal, "\n".join(buf), False)

            status = "Dibatalkan" if self.stop.is_set() else "Selesai"
            self.post(self.on_done, baru, upd, dup, gagal, status)
        except Exception as ex:
            self.post(self.on_progress, processed, max(total, 1), baru, upd, dup, gagal,
                      f"[ERROR] Scan berhenti: {ex}", False)
            self.post(self.on_done, baru, upd, dup, gagal, "Berhenti (error)")
        finally:
            if db is not None:
                db.close()

    def on_progress(self, i, total, baru, upd, dup, gagal, text, refresh):
        if self.stop.is_set():
            self.progress.set(i * 100.0 / total if total else 0, "Membatalkan...")
        else:
            self.progress.set(i * 100.0 / total if total else 0, f"Memproses {i}/{total}")
        self.lbl_stat.SetLabel(self.stat_text(baru, upd, dup, gagal))
        if text:
            self.log.AppendText(text + "\n")
        if refresh:
            self.refresh_grid()

    def on_done(self, baru, upd, dup, gagal, status):
        self.lbl_stat.SetLabel(self.stat_text(baru, upd, dup, gagal))
        if status == "Selesai":
            self.progress.set(100, "Selesai")
        else:
            self.progress.set(self.progress.value, status)
        self.log.AppendText(f"--- {status} ---\n")
        self.btn_folder.Enable()
        self.btn_merge.Enable()
        self.btn_cancel.Disable()
        self.refresh_grid()

    def on_close(self, e):
        self.stop.set()
        if self.search_timer is not None:
            self.search_timer.Stop()
        if self.worker and self.worker.is_alive():
            self.worker.join(timeout=2)  # beri waktu singkat agar thread berhenti rapi
        self.db.close()
        e.Skip()


if __name__ == "__main__":
    app = wx.App(False)
    frame = MainFrame()
    frame.Show()
    app.MainLoop()
