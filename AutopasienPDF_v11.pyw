# -*- coding: utf-8 -*-
"""
Auto List Pasien dari PDF
Database : AutopasienPDF.db (SQLite, bisa dibuka pakai DBeaver)
Library  : wxPython, pymupdf, openpyxl
           pip install wxPython pymupdf openpyxl
"""
import json
import os
import re
import sqlite3
import sys
import threading
import webbrowser
from datetime import datetime

import wx
import wx.grid as gridlib

try:
    import pymupdf as fitz  # nama baru
except ImportError:
    try:
        import fitz  # nama lama
    except ImportError:
        fitz = None
try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None
try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None


def get_app_dir():
    """Folder aplikasi: folder .exe kalau sudah di-convert, folder script kalau jalan biasa."""
    if getattr(sys, "frozen", False):
        return os.path.dirname(os.path.abspath(sys.executable))
    return os.path.dirname(os.path.abspath(__file__))


APP_TITLE = "Auto List Pasien dari pdf"
DB_NAME = "AutopasienPDF.db"
BASE_DIR = get_app_dir()
DEFAULT_DB = os.path.join(BASE_DIR, DB_NAME)
CONFIG_FILE = os.path.join(BASE_DIR, "AutoListPasien.json")

GENDERS = ["Laki-laki", "Perempuan"]

# (nama kolom db, label tampilan)
COLS = [
    ("id", "ID"),
    ("nama_peserta", "NAMA PESERTA"),
    ("no_peserta", "NO. PESERTA"),
    ("no_rekammedis", "NO REKAM MEDIS"),
    ("umur", "UMUR"),
    ("tanggal_lahir", "TGL LAHIR"),
    ("no_SEP", "NO SEP"),
    ("jenis_kelamin", "JENIS KELAMIN"),
    ("tgl_masuk", "TGL MASUK"),
    ("tgl_keluar", "TGL KELUAR"),
    ("no_telepon", "NO TELEPON"),
]
DB_FIELDS = [c[0] for c in COLS if c[0] != "id"]
PHONE_COL = 10
WA_COL = 11
COL_WIDTHS = [55, 230, 140, 120, 60, 100, 240, 120, 100, 100, 130, 110]

NUMERIC_COLS = {"id", "umur"}
DATE_COLS = {"tanggal_lahir", "tgl_masuk", "tgl_keluar"}

COLOR_WHITE = wx.Colour(255, 255, 255)
COLOR_GREY = wx.Colour(232, 232, 232)
COLOR_WA = wx.Colour(37, 211, 102)


# ------------------------------------------------------ simpan pilihan database
def load_last_db():
    """Ambil database terakhir yang dipilih user (kalau file-nya masih ada)."""
    try:
        with open(CONFIG_FILE, "r", encoding="utf-8") as f:
            path = json.load(f).get("db_path", "")
        if path and os.path.isfile(path):
            return path
    except Exception:
        pass
    return DEFAULT_DB


def save_last_db(path):
    try:
        with open(CONFIG_FILE, "w", encoding="utf-8") as f:
            json.dump({"db_path": path}, f, ensure_ascii=False, indent=2)
    except Exception:
        pass


# ---------------------------------------------------------------- PDF parsing
FIELD_LABELS = {
    "no_peserta": ["Nomor Peserta", "No. Peserta", "No Peserta"],
    "no_rekammedis": ["Nomor Rekam Medis", "No. Rekam Medis", "No Rekam Medis", "No. RM"],
    "umur": ["Umur Tahun", "Umur"],
    "tanggal_lahir": ["Tanggal Lahir", "Tgl. Lahir", "Tgl Lahir"],
    "no_SEP": ["Nomor SEP", "No. SEP", "No SEP"],
    "tgl_masuk": ["Tanggal Masuk", "Tgl. Masuk", "Tgl Masuk"],
    "tgl_keluar": ["Tanggal Keluar", "Tgl. Keluar", "Tgl Keluar"],
}
NAME_LABELS = ["Nama Peserta", "Nama Pasien", "Nama"]
GENDER_LABELS = ["Jenis Kelamin"]

# "No. Telepon : 083134676447" (nilai boleh turun ke baris bawah label)
PHONE_LABEL_RE = re.compile(
    r"No\.?\s*(?:Telepon|Telpon|Telp|Tlp|HP|Handphone)\s*[:：]?\s*(\+?\d[\d\-\. ]{7,22}\d)",
    re.IGNORECASE,
)
# cadangan: nomor HP berpola 08xxxxxxxxxx / 628xxxxxxxxxx / +628xxxxxxxxxx
PHONE_LIKE_RE = re.compile(r"(?<![\d])((?:\+?62|0)8\d{8,12})(?![\d])")


class PdfDoc:
    def __init__(self, path):
        self.doc = None
        self.reader = None
        if fitz:
            self.doc = fitz.open(path)
            self.count = self.doc.page_count
        elif PdfReader:
            self.reader = PdfReader(path)
            self.count = len(self.reader.pages)
        else:
            raise RuntimeError("Install dulu: pip install pymupdf")
        self._cache = {}

    def text(self, i):
        if i in self._cache:
            return self._cache[i]
        try:
            if self.doc is not None:
                t = self.doc[i].get_text("text") or ""
            else:
                t = self.reader.pages[i].extract_text() or ""
        except Exception:
            t = ""
        self._cache[i] = t
        return t

    def close(self):
        if self.doc is not None:
            self.doc.close()


def find_field(text, labels):
    """Cari 'Label : nilai' di teks. Kalau nilai kosong, ambil baris berikutnya."""
    for lab in labels:
        words = lab.split()
        pat = r"(?<![A-Za-z])" + r"\s*".join(re.escape(w) for w in words) + r"\s*[:：]\s*([^\n]*)"
        m = re.search(pat, text, re.IGNORECASE)
        if not m:
            continue
        val = m.group(1).strip()
        if not val:
            rest = text[m.end():].split("\n")
            for line in rest:
                line = line.strip()
                if line:
                    if ":" not in line:
                        val = line
                    break
        val = re.split(r"\s{2,}|\t", val)[0].strip()
        if val:
            return val
    return ""


def normalize_gender(val):
    low = (val or "").lower()
    if "laki" in low:
        return "Laki-laki"
    if "perempuan" in low or "wanita" in low:
        return "Perempuan"
    return ""


def first_number(val):
    m = re.search(r"\d+", val or "")
    return m.group(0) if m else ""


def clean_phone(val):
    """Ambil angka saja: '0831-3467-6447' -> '083134676447'."""
    return re.sub(r"\D", "", (val or "").strip())


def phone_from_text(text, exclude):
    """Cari nomor telepon di satu halaman. exclude = nomor lain yang bukan telepon."""
    for m in PHONE_LABEL_RE.finditer(text):
        num = clean_phone(m.group(1))
        if 9 <= len(num) <= 15 and num not in exclude:
            return num
    for m in PHONE_LIKE_RE.finditer(text):
        num = clean_phone(m.group(1))
        if num not in exclude:
            return num
    return ""


def extract_phone(pdf, exclude):
    """
    Urutan pencarian:
    1. halaman yang ada tulisan SURAT ELEGIBILITAS PESERTA
    2. halaman ke-2
    3. halaman ke-1
    4. semua halaman
    """
    order = []
    for i in range(pdf.count):
        t = pdf.text(i).upper()
        if "ELEGIBILITAS" in t or "ELIGIBILITAS" in t:
            order.append(i)
    if pdf.count >= 2 and 1 not in order:
        order.append(1)
    if 0 not in order:
        order.append(0)
    for i in range(pdf.count):
        if i not in order:
            order.append(i)

    for i in order:
        num = phone_from_text(pdf.text(i), exclude)
        if num:
            return num
    return ""


def extract_patient(path):
    pdf = PdfDoc(path)
    try:
        if pdf.count == 0:
            return None
        p1 = pdf.text(0)
        data = {k: find_field(p1, labs) for k, labs in FIELD_LABELS.items()}
        data["umur"] = first_number(data["umur"])

        # Jenis kelamin (halaman 1 dulu, kalau tidak ada cari semua halaman)
        gender = normalize_gender(find_field(p1, GENDER_LABELS))
        if not gender:
            for i in range(pdf.count):
                gender = normalize_gender(find_field(pdf.text(i), GENDER_LABELS))
                if gender:
                    break
        data["jenis_kelamin"] = gender

        # Nama: halaman 2 dulu, kalau tidak ketemu cari semua halaman
        nama = ""
        if pdf.count >= 2:
            nama = find_field(pdf.text(1), NAME_LABELS)
        if not nama:
            for i in range(pdf.count):
                nama = find_field(pdf.text(i), NAME_LABELS)
                if nama:
                    break
        data["nama_peserta"] = nama

        # No. Telepon: halaman "SURAT ELEGIBILITAS PESERTA" (halaman 2)
        exclude = set()
        for k in ("no_peserta", "no_SEP", "no_rekammedis"):
            v = re.sub(r"\D", "", data.get(k, "") or "")
            if v:
                exclude.add(v)
        data["no_telepon"] = extract_phone(pdf, exclude)
        return data
    finally:
        pdf.close()


def row_get(row, key):
    """Ambil nilai dari sqlite3.Row dengan aman (kolom tidak ada -> kosong)."""
    try:
        if key in row.keys():
            val = row[key]
            return "" if val is None else val
    except Exception:
        pass
    return ""


def sort_key(col_key, value):
    """Kunci pengurutan: angka sebagai angka, tanggal sebagai tanggal, sisanya teks."""
    text = str(value).strip()
    if col_key in NUMERIC_COLS:
        try:
            return (0, float(text))
        except ValueError:
            return (1, text.lower())
    if col_key in DATE_COLS:
        m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})", text)
        if m:
            d, mo, y = m.groups()
            return (0, (int(y), int(mo), int(d)))
        return (1, text.lower())
    return (0, text.lower()) if text else (1, "")


# ------------------------------------------------------------------- Database
class Database:
    def __init__(self, path):
        self.path = path
        self.conn = sqlite3.connect(path)
        self.conn.row_factory = sqlite3.Row
        self.create()

    def create(self):
        self.conn.execute(
            """CREATE TABLE IF NOT EXISTS pasien (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                nama_peserta TEXT,
                no_peserta TEXT,
                no_rekammedis TEXT,
                umur TEXT,
                tanggal_lahir TEXT,
                no_SEP TEXT,
                jenis_kelamin TEXT,
                tgl_masuk TEXT,
                tgl_keluar TEXT,
                no_telepon TEXT
            )"""
        )
        self.conn.commit()
        self.migrate()

    def migrate(self):
        """Samakan tabel lama (kolom versi sebelumnya) dengan struktur terbaru."""
        cols = {r[1].lower() for r in self.conn.execute("PRAGMA table_info(pasien)")}

        if "nama_pasien" in cols and "nama_peserta" not in cols:
            self.conn.execute("ALTER TABLE pasien RENAME COLUMN Nama_Pasien TO nama_peserta")
            cols.discard("nama_pasien")
            cols.add("nama_peserta")

        for field in DB_FIELDS:
            if field.lower() not in cols:
                self.conn.execute(f"ALTER TABLE pasien ADD COLUMN {field} TEXT")
        self.conn.commit()

    def count(self):
        return self.conn.execute("SELECT COUNT(*) FROM pasien").fetchone()[0]

    def find_existing(self, d):
        """Kembalikan baris yang sama (atau None)."""
        if d.get("no_SEP"):
            cur = self.conn.execute("SELECT * FROM pasien WHERE no_SEP=?", (d["no_SEP"],))
        else:
            cur = self.conn.execute(
                "SELECT * FROM pasien WHERE nama_peserta=? AND no_peserta=? AND tgl_masuk=?",
                (d.get("nama_peserta", ""), d.get("no_peserta", ""), d.get("tgl_masuk", "")),
            )
        return cur.fetchone()

    def update_phone(self, pid, phone):
        self.conn.execute("UPDATE pasien SET no_telepon=? WHERE id=?", (phone, pid))
        self.conn.commit()

    def insert(self, d):
        cols = ",".join(DB_FIELDS)
        qs = ",".join("?" * len(DB_FIELDS))
        self.conn.execute(
            f"INSERT INTO pasien ({cols}) VALUES ({qs})",
            [d.get(k, "") for k in DB_FIELDS],
        )
        self.conn.commit()

    def fetch(self, keyword=""):
        if keyword:
            like = f"%{keyword}%"
            cur = self.conn.execute(
                """SELECT * FROM pasien WHERE nama_peserta LIKE ? OR no_peserta LIKE ?
                   OR no_rekammedis LIKE ? OR no_SEP LIKE ? OR no_telepon LIKE ?
                   OR tgl_masuk LIKE ? OR tgl_keluar LIKE ? OR tanggal_lahir LIKE ?
                   OR jenis_kelamin LIKE ? OR umur LIKE ? ORDER BY id DESC""",
                (like,) * 10,
            )
        else:
            cur = self.conn.execute("SELECT * FROM pasien ORDER BY id DESC")
        return cur.fetchall()

    def get(self, pid):
        return self.conn.execute("SELECT * FROM pasien WHERE id=?", (pid,)).fetchone()

    def update(self, pid, d):
        sets = ",".join(f"{k}=?" for k in DB_FIELDS)
        self.conn.execute(
            f"UPDATE pasien SET {sets} WHERE id=?",
            [d.get(k, "") for k in DB_FIELDS] + [pid],
        )
        self.conn.commit()

    def delete(self, pid):
        self.conn.execute("DELETE FROM pasien WHERE id=?", (pid,))
        self.conn.commit()

    def close(self):
        try:
            self.conn.close()
        except Exception:
            pass


# ------------------------------------------------------------------ Progress
class ProgressBar(wx.Panel):
    """Loading bar warna hijau + persen."""

    def __init__(self, parent):
        super().__init__(parent, size=(-1, 30))
        self.SetMinSize((-1, 30))
        self.SetBackgroundStyle(wx.BG_STYLE_PAINT)
        self.value = 0.0
        self.text = "Siap"
        self.Bind(wx.EVT_PAINT, self.on_paint)
        self.Bind(wx.EVT_SIZE, self.on_size)

    def on_size(self, e):
        self.Refresh()
        e.Skip()

    def set(self, value, text=None):
        self.value = max(0.0, min(100.0, value))
        if text is not None:
            self.text = text
        self.Refresh()
        self.Update()

    def on_paint(self, _):
        dc = wx.AutoBufferedPaintDC(self)
        w, h = self.GetClientSize()
        dc.SetBackground(wx.Brush(wx.Colour(228, 228, 228)))
        dc.Clear()
        dc.SetPen(wx.TRANSPARENT_PEN)
        dc.SetBrush(wx.Brush(wx.Colour(46, 184, 71)))
        dc.DrawRectangle(0, 0, int(w * self.value / 100.0), h)
        font = self.GetFont()
        font.SetWeight(wx.FONTWEIGHT_BOLD)
        dc.SetFont(font)
        dc.SetTextForeground(wx.Colour(0, 0, 0))
        label = f"{self.text}  {self.value:.0f}%"
        tw, th = dc.GetTextExtent(label)
        dc.DrawText(label, (w - tw) // 2, (h - th) // 2)


# --------------------------------------------------------------- Dialog Ubah
class EditDialog(wx.Dialog):
    def __init__(self, parent, row):
        super().__init__(parent, title="Ubah Data Pasien", style=wx.DEFAULT_DIALOG_STYLE)
        self.ctrls = {}
        grid = wx.FlexGridSizer(0, 2, 8, 10)
        grid.AddGrowableCol(1)
        for key, label in COLS:
            if key == "id":
                continue
            val = row_get(row, key)
            grid.Add(wx.StaticText(self, label=label.title()), 0, wx.ALIGN_CENTER_VERTICAL)
            if key == "jenis_kelamin":
                ctrl = wx.Choice(self, choices=GENDERS)
                if val in GENDERS:
                    ctrl.SetSelection(GENDERS.index(val))
            else:
                ctrl = wx.TextCtrl(self, value=str(val), size=(300, -1))
            self.ctrls[key] = ctrl
            grid.Add(ctrl, 1, wx.EXPAND)

        main = wx.BoxSizer(wx.VERTICAL)
        main.Add(grid, 1, wx.ALL | wx.EXPAND, 15)
        main.Add(self.CreateStdDialogButtonSizer(wx.OK | wx.CANCEL), 0, wx.ALL | wx.ALIGN_RIGHT, 10)
        self.SetSizerAndFit(main)
        self.CentreOnParent()

    def values(self):
        out = {}
        for key, ctrl in self.ctrls.items():
            if isinstance(ctrl, wx.Choice):
                idx = ctrl.GetSelection()
                out[key] = GENDERS[idx] if idx != wx.NOT_FOUND else ""
            else:
                out[key] = ctrl.GetValue().strip()
        return out


# ---------------------------------------------------------------- Main Frame
class MainFrame(wx.Frame):
    def __init__(self):
        super().__init__(None, title=APP_TITLE, size=(1350, 720))
        self.cancel = False
        self.worker = None
        self.sort_col = None   # index kolom yang sedang di-sort
        self.sort_asc = True
        self.search_timer = None
        self.current_rows = []  # baris yang sedang tampil di list (dipakai export)

        # buka database terakhir yang dipilih (atau default di folder .exe)
        self.db_path = load_last_db()
        try:
            self.db = Database(self.db_path)
        except Exception as ex:
            wx.MessageBox(
                f"Gagal membuka database:\n{self.db_path}\n\n{ex}\n\nMemakai database default.",
                "Error", wx.OK | wx.ICON_ERROR,
            )
            self.db_path = DEFAULT_DB
            self.db = Database(self.db_path)
        self.SetTitle(f"{APP_TITLE} - {self.db_path}")

        main = wx.BoxSizer(wx.HORIZONTAL)
        self.left = self.build_left()
        self.right = self.build_right()
        main.Add(self.left, 2, wx.EXPAND)   # 20%
        main.Add(self.right, 8, wx.EXPAND)  # 80%
        self.SetSizer(main)

        self.Bind(wx.EVT_CLOSE, self.on_close)
        self.Centre()
        self.Maximize(True)  # otomatis maximize, tetap bisa minimize / restore down
        self.refresh_grid()

    # ---------------- sisi kiri (20%)
    def build_left(self):
        p = wx.Panel(self)
        p.SetBackgroundColour(wx.Colour(245, 247, 250))
        s = wx.BoxSizer(wx.VERTICAL)

        title = wx.StaticText(p, label="SELEKSI FOLDER PDF")
        f = title.GetFont()
        f.SetWeight(wx.FONTWEIGHT_BOLD)
        f.SetPointSize(f.GetPointSize() + 2)
        title.SetFont(f)
        s.Add(title, 0, wx.ALL, 12)

        self.btn_folder = wx.Button(p, label="PILIH FOLDER (+ SUBFOLDER)", size=(-1, 40))
        self.btn_folder.Bind(wx.EVT_BUTTON, self.on_pick_folder)
        s.Add(self.btn_folder, 0, wx.LEFT | wx.RIGHT | wx.EXPAND, 12)

        self.btn_cancel = wx.Button(p, label="BATALKAN PROSES")
        self.btn_cancel.Disable()
        self.btn_cancel.Bind(wx.EVT_BUTTON, lambda e: setattr(self, "cancel", True))
        s.Add(self.btn_cancel, 0, wx.LEFT | wx.RIGHT | wx.TOP | wx.EXPAND, 12)

        self.lbl_folder = wx.StaticText(p, label="Folder: -", style=wx.ST_ELLIPSIZE_MIDDLE)
        s.Add(self.lbl_folder, 0, wx.ALL | wx.EXPAND, 12)

        self.lbl_stat = wx.StaticText(p, label=self.stat_text(0, 0, 0, 0))
        s.Add(self.lbl_stat, 0, wx.LEFT | wx.RIGHT | wx.BOTTOM, 12)

        self.log = wx.TextCtrl(p, style=wx.TE_MULTILINE | wx.TE_READONLY | wx.HSCROLL)
        s.Add(self.log, 1, wx.LEFT | wx.RIGHT | wx.BOTTOM | wx.EXPAND, 12)

        p.SetSizer(s)
        return p

    @staticmethod
    def stat_text(baru, upd, dup, gagal):
        return f"Baru: {baru}   Update: {upd}\nDuplikat: {dup}   Gagal: {gagal}"

    # ---------------- sisi kanan (80%)
    def build_right(self):
        p = wx.Panel(self)
        s = wx.BoxSizer(wx.VERTICAL)

        bar = wx.BoxSizer(wx.HORIZONTAL)
        self.btn_db = wx.Button(p, label="SELEKSI DATABASE")
        self.btn_edit = wx.Button(p, label="UBAH DATA")
        self.btn_del = wx.Button(p, label="HAPUS")
        self.btn_ref = wx.Button(p, label="REFRESH")
        self.btn_export = wx.Button(p, label="EXPORT TO EXCEL")
        self.btn_export.SetBackgroundColour(wx.Colour(33, 115, 70))
        self.btn_export.SetForegroundColour(wx.Colour(255, 255, 255))
        self.txt_search = wx.TextCtrl(p, size=(300, -1))
        self.txt_search.SetHint("Ketik untuk mencari (nama / no. peserta / SEP / telp...)")
        self.lbl_total = wx.StaticText(p, label="Total: 0")

        for b in (self.btn_db, self.btn_edit, self.btn_del, self.btn_ref, self.btn_export):
            bar.Add(b, 0, wx.RIGHT, 6)
        bar.AddStretchSpacer(1)
        bar.Add(self.lbl_total, 0, wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 10)
        bar.Add(self.txt_search, 0)
        s.Add(bar, 0, wx.ALL | wx.EXPAND, 8)

        self.grid = gridlib.Grid(p)
        self.grid.CreateGrid(0, len(COLS) + 1)
        for i, w in enumerate(COL_WIDTHS):
            self.grid.SetColSize(i, w)
        self.grid.SetRowLabelSize(0)
        self.grid.EnableEditing(False)
        self.grid.SetSelectionMode(getattr(gridlib.Grid, "GridSelectRows", getattr(gridlib.Grid, "SelectRows", 1)))
        self.update_header_labels()

        attr = gridlib.GridCellAttr()
        attr.SetBackgroundColour(COLOR_WA)
        attr.SetTextColour(wx.Colour(255, 255, 255))
        attr.SetAlignment(wx.ALIGN_CENTRE, wx.ALIGN_CENTRE)
        fnt = wx.Font(self.grid.GetDefaultCellFont())
        fnt.SetWeight(wx.FONTWEIGHT_BOLD)
        attr.SetFont(fnt)
        self.grid.SetColAttr(WA_COL, attr)

        s.Add(self.grid, 1, wx.LEFT | wx.RIGHT | wx.EXPAND, 8)

        self.progress = ProgressBar(p)
        s.Add(self.progress, 0, wx.ALL | wx.EXPAND, 8)

        self.btn_db.Bind(wx.EVT_BUTTON, self.on_select_db)
        self.btn_edit.Bind(wx.EVT_BUTTON, self.on_edit)
        self.btn_del.Bind(wx.EVT_BUTTON, self.on_delete)
        self.btn_ref.Bind(wx.EVT_BUTTON, lambda e: self.refresh_grid())
        self.btn_export.Bind(wx.EVT_BUTTON, self.on_export_excel)
        self.txt_search.Bind(wx.EVT_TEXT, self.on_search_text)
        self.grid.Bind(gridlib.EVT_GRID_CELL_LEFT_CLICK, self.on_cell_click)
        self.grid.Bind(gridlib.EVT_GRID_CELL_LEFT_DCLICK, self.on_cell_dclick)
        self.grid.Bind(gridlib.EVT_GRID_LABEL_LEFT_CLICK, self.on_header_click)

        p.SetSizer(s)
        return p

    # ---------------- pencarian langsung
    def on_search_text(self, _):
        if self.search_timer is not None:
            self.search_timer.Stop()
        self.search_timer = wx.CallLater(150, self.refresh_grid)

    # ---------------- sort header
    def update_header_labels(self):
        for i, (_, label) in enumerate(COLS):
            arrow = ""
            if i == self.sort_col:
                arrow = "  ▲" if self.sort_asc else "  ▼"
            self.grid.SetColLabelValue(i, label + arrow)
        self.grid.SetColLabelValue(WA_COL, "WHATSAPP")

    def on_header_click(self, e):
        col = e.GetCol()
        if col < 0 or col >= len(COLS):
            return
        if self.sort_col == col:
            self.sort_asc = not self.sort_asc
        else:
            self.sort_col = col
            self.sort_asc = True
        self.update_header_labels()
        self.refresh_grid()

    def sort_rows(self, rows):
        if self.sort_col is None:
            return list(rows)
        key = COLS[self.sort_col][0]
        return sorted(rows, key=lambda r: sort_key(key, row_get(r, key)), reverse=not self.sort_asc)

    # ---------------- grid
    def refresh_grid(self):
        try:
            rows = self.sort_rows(self.db.fetch(self.txt_search.GetValue().strip()))
        except Exception as ex:
            wx.MessageBox(f"Gagal membaca database:\n{ex}", "Error", wx.OK | wx.ICON_ERROR)
            return
        self.current_rows = rows  # simpan baris yang tampil, dipakai untuk export
        g = self.grid
        g.BeginBatch()
        n = g.GetNumberRows()
        if n:
            g.DeleteRows(0, n)
        g.AppendRows(len(rows))
        white_text = wx.Colour(255, 255, 255)
        for r, row in enumerate(rows):
            bg = COLOR_WHITE if r % 2 == 0 else COLOR_GREY  # baris belang
            for c, (key, _) in enumerate(COLS):
                g.SetCellValue(r, c, str(row_get(row, key)))
                g.SetCellBackgroundColour(r, c, bg)
            g.SetCellValue(r, WA_COL, "WhatsApp")
            g.SetCellBackgroundColour(r, WA_COL, COLOR_WA)
            g.SetCellTextColour(r, WA_COL, white_text)
        g.EndBatch()
        self.lbl_total.SetLabel(f"Total: {len(rows)}")
        self.right.Layout()

    def selected_id(self):
        g = self.grid
        rows = g.GetSelectedRows()
        row = rows[0] if rows else g.GetGridCursorRow()
        if row is None or row < 0 or row >= g.GetNumberRows():
            return None
        try:
            return int(g.GetCellValue(row, 0))
        except ValueError:
            return None

    def on_cell_click(self, e):
        row, col = e.GetRow(), e.GetCol()
        if col == WA_COL:
            self.open_whatsapp(self.grid.GetCellValue(row, PHONE_COL))
        e.Skip()

    def on_cell_dclick(self, e):
        """Double klik baris -> UBAH DATA (kolom WhatsApp diabaikan)."""
        row, col = e.GetRow(), e.GetCol()
        if col == WA_COL:
            return
        try:
            pid = int(self.grid.GetCellValue(row, 0))
        except ValueError:
            return
        self.grid.SelectRow(row)
        self.edit_by_id(pid)

    def open_whatsapp(self, phone):
        digits = re.sub(r"\D", "", phone or "")
        if not digits:
            wx.MessageBox("Nomor telepon kosong.", "WhatsApp", wx.OK | wx.ICON_WARNING)
            return
        if digits.startswith("0"):
            digits = "62" + digits[1:]
        elif digits.startswith("8"):
            digits = "62" + digits
        webbrowser.open(f"https://wa.me/{digits}")

    # ---------------- export excel
    def on_export_excel(self, _):
        """Export data yang sedang tampil di list (sesuai pencarian & urutan)."""
        if openpyxl is None:
            wx.MessageBox(
                "Library Excel belum terpasang.\n\nJalankan di CMD:\npip install openpyxl",
                "Export to Excel", wx.OK | wx.ICON_WARNING,
            )
            return
        rows = self.current_rows
        if not rows:
            wx.MessageBox("Tidak ada data untuk diexport.", "Export to Excel", wx.OK | wx.ICON_INFORMATION)
            return

        keyword = self.txt_search.GetValue().strip()
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        default_name = f"Data_Pasien_{'Hasil_Pencarian_' if keyword else ''}{stamp}.xlsx"

        dlg = wx.FileDialog(
            self, "Simpan sebagai Excel", defaultDir=BASE_DIR, defaultFile=default_name,
            wildcard="Excel (*.xlsx)|*.xlsx", style=wx.FD_SAVE | wx.FD_OVERWRITE_PROMPT,
        )
        if dlg.ShowModal() != wx.ID_OK:
            dlg.Destroy()
            return
        path = dlg.GetPath()
        dlg.Destroy()
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Data Pasien"

            # header
            ws.append([label for _, label in COLS])
            head_fill = PatternFill("solid", fgColor="2EB847")
            for c in range(1, len(COLS) + 1):
                cell = ws.cell(row=1, column=c)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = head_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # isi data (semua disimpan sebagai teks supaya angka 0 di depan tidak hilang)
            for row in rows:
                ws.append([str(row_get(row, key)) for key, _ in COLS])
            for r in range(2, len(rows) + 2):
                for c in range(1, len(COLS) + 1):
                    ws.cell(row=r, column=c).number_format = "@"

            # lebar kolom otomatis
            for c, (key, label) in enumerate(COLS, 1):
                longest = len(label)
                for row in rows:
                    longest = max(longest, len(str(row_get(row, key))))
                ws.column_dimensions[get_column_letter(c)].width = min(longest + 3, 45)

            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            wb.save(path)
        except PermissionError:
            wx.MessageBox(
                "File tidak bisa disimpan. Kemungkinan file itu sedang terbuka di Excel.\n"
                "Tutup dulu file-nya atau pakai nama lain.",
                "Export to Excel", wx.OK | wx.ICON_ERROR,
            )
            return
        except Exception as ex:
            wx.MessageBox(f"Gagal export:\n{ex}", "Export to Excel", wx.OK | wx.ICON_ERROR)
            return

        self.progress.set(100, f"Export selesai: {len(rows)} data")
        ask = wx.MessageBox(
            f"Berhasil export {len(rows)} data ke:\n{path}\n\nBuka file sekarang?",
            "Export to Excel", wx.YES_NO | wx.ICON_INFORMATION,
        )
        if ask == wx.YES:
            try:
                os.startfile(path)  # Windows
            except Exception:
                pass

    # ---------------- tombol
    def on_select_db(self, _):
        start_dir = os.path.dirname(self.db_path) if os.path.isdir(os.path.dirname(self.db_path)) else BASE_DIR
        dlg = wx.FileDialog(
            self, "Pilih / buat database", defaultDir=start_dir,
            defaultFile=os.path.basename(self.db_path),
            wildcard="SQLite DB (*.db)|*.db|Semua file (*.*)|*.*", style=wx.FD_SAVE,
        )
        if dlg.ShowModal() == wx.ID_OK:
            path = dlg.GetPath()
            if not os.path.exists(path) and not os.path.splitext(path)[1]:
                path += ".db"
            try:
                new_db = Database(path)
                total = new_db.count()  # tes baca tabel
            except Exception as ex:
                wx.MessageBox(f"Gagal membuka database:\n{path}\n\n{ex}", "Error", wx.OK | wx.ICON_ERROR)
                dlg.Destroy()
                return
            self.db.close()
            self.db = new_db
            self.db_path = path
            save_last_db(path)  # diingat untuk pembukaan berikutnya
            self.SetTitle(f"{APP_TITLE} - {path}")
            self.refresh_grid()
            self.progress.set(0, f"Database dibuka: {total} data")
        dlg.Destroy()

    def edit_by_id(self, pid):
        row = self.db.get(pid)
        if not row:
            return
        dlg = EditDialog(self, row)
        if dlg.ShowModal() == wx.ID_OK:
            self.db.update(pid, dlg.values())
            self.refresh_grid()
        dlg.Destroy()

    def on_edit(self, _):
        pid = self.selected_id()
        if pid is None:
            wx.MessageBox("Klik salah satu data di list dulu.", "Ubah Data", wx.OK | wx.ICON_INFORMATION)
            return
        self.edit_by_id(pid)

    def on_delete(self, _):
        pid = self.selected_id()
        if pid is None:
            wx.MessageBox("Klik salah satu data di list dulu.", "Hapus", wx.OK | wx.ICON_INFORMATION)
            return
        if wx.MessageBox("Yakin hapus data ini?", "Konfirmasi", wx.YES_NO | wx.ICON_QUESTION) == wx.YES:
            self.db.delete(pid)
            self.refresh_grid()

    # ---------------- scan folder
    def on_pick_folder(self, _):
        if self.worker and self.worker.is_alive():
            return
        dlg = wx.DirDialog(self, "Pilih folder PDF (termasuk subfolder)", style=wx.DD_DEFAULT_STYLE)
        if dlg.ShowModal() == wx.ID_OK:
            folder = dlg.GetPath()
            self.lbl_folder.SetLabel(f"Folder: {folder}")
            self.lbl_folder.SetToolTip(folder)
            self.log.Clear()
            self.cancel = False
            self.btn_folder.Disable()
            self.btn_cancel.Enable()
            self.progress.set(0, "Memulai...")
            self.worker = threading.Thread(target=self.scan_worker, args=(folder,), daemon=True)
            self.worker.start()
        dlg.Destroy()

    def scan_worker(self, folder):
        db = Database(self.db_path)  # koneksi sendiri untuk thread ini
        files = []
        for root, _, names in os.walk(folder):
            for n in names:
                if n.lower().endswith(".pdf"):
                    files.append(os.path.join(root, n))
        files.sort()
        total = len(files)
        baru = upd = dup = gagal = 0

        if total == 0:
            wx.CallAfter(self.on_done, 0, 0, 0, 0, "Tidak ada file PDF.")
            db.close()
            return

        for i, path in enumerate(files, 1):
            if self.cancel:
                break
            name = os.path.basename(path)
            try:
                data = extract_patient(path)
                if not data or not (data["nama_peserta"] or data["no_peserta"] or data["no_SEP"]):
                    gagal += 1
                    msg = f"[GAGAL] {name} - data tidak ditemukan"
                else:
                    existing = db.find_existing(data)
                    if existing is None:
                        db.insert(data)
                        baru += 1
                        msg = f"[BARU ] {name} - {data['nama_peserta']} | telp: {data['no_telepon'] or '-'}"
                    else:
                        old_phone = str(row_get(existing, "no_telepon")).strip()
                        if not old_phone and data["no_telepon"]:
                            db.update_phone(existing["id"], data["no_telepon"])
                            upd += 1
                            msg = f"[UPDATE] {name} - telp diisi: {data['no_telepon']}"
                        else:
                            dup += 1
                            msg = f"[SKIP ] {name} - sudah ada ({data['nama_peserta']})"
            except Exception as ex:
                gagal += 1
                msg = f"[ERROR] {name} - {ex}"
            wx.CallAfter(self.on_progress, i, total, baru, upd, dup, gagal, msg)

        db.close()
        status = "Dibatalkan" if self.cancel else "Selesai"
        wx.CallAfter(self.on_done, baru, upd, dup, gagal, status)

    def on_progress(self, i, total, baru, upd, dup, gagal, msg):
        self.progress.set(i * 100.0 / total, f"Memproses {i}/{total}")
        self.lbl_stat.SetLabel(self.stat_text(baru, upd, dup, gagal))
        self.log.AppendText(msg + "\n")
        if i % 25 == 0:
            self.refresh_grid()

    def on_done(self, baru, upd, dup, gagal, status):
        self.lbl_stat.SetLabel(self.stat_text(baru, upd, dup, gagal))
        if status == "Selesai":
            self.progress.set(100, "Selesai")
        else:
            self.progress.set(self.progress.value, status)
        self.log.AppendText(f"--- {status} ---\n")
        self.btn_folder.Enable()
        self.btn_cancel.Disable()
        self.refresh_grid()

    def on_close(self, e):
        self.cancel = True
        if self.search_timer is not None:
            self.search_timer.Stop()
        self.db.close()
        e.Skip()


if __name__ == "__main__":
    app = wx.App(False)
    frame = MainFrame()
    frame.Show()
    app.MainLoop()
